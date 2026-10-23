#!/usr/bin/env python3
"""
Gera a visao de gestores por level e a aba de divergencias.

Saidas:
  estrutura/Gestores_por_Level.md     - um bloco por level, quem administra quem
  estrutura/Gestores_por_Level.xlsx   - abas anteriores + aba 'Divergencias'

Uso:  python3 estrutura/gestores_por_level.py
"""
import os
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cross_check import (toks, CARGOS_DRAFT, C_LEVEL_DRE, DRAFT_LIDERADOS,
                         MOVIDOS_PARA_CEO, NOTAS_DRAFT, PAT_INTERN, RAIZ,
                         RISCADOS, SEM_CARGO, TESTE, confere_draft,
                         diretoria_de, divergencias, ler_aba, ler_factorial,
                         montar)

SAIDA = os.path.join(RAIZ, "estrutura")


def plural(n, s, p=None):
    return f"{n} {s}" if n == 1 else f"{n} {p or s + 's'}"


def cargo_de(p):
    """Cargo da planilha; para L2/L3 usa o cargo lido no draft."""
    if p["nome"] in CARGOS_DRAFT:
        cargo, slide = CARGOS_DRAFT[p["nome"]]
        return cargo, f"draft {slide}"
    return (p["cargo"], "planilha") if p["cargo"] else (SEM_CARGO, "—")


def confere(p, como_esta):
    """Texto de conferencia contra o numero de liderados impresso no draft."""
    d = DRAFT_LIDERADOS.get(p["nome"])
    if not d:
        return "não consta no draft"
    base = como_esta.get(p["nome"], p["total"])
    if d[0] == base:
        return f"confere ({d[1]})"
    return f"draft {d[0]}, planilha {base} (+{base - d[0]} admitido depois)"


# ------------------------------------------------------------------ markdown
def gerar_md(todos, filhos, ceo, como_esta, dados):
    por_level = defaultdict(list)
    for p in todos:
        por_level[p["level"]].append(p)
    area = diretoria_de(todos, filhos, ceo)

    L = []
    A = L.append
    A("# Gestores por level — quem administra quem")
    A("")
    A("Fontes: **`Colaboradores-factorial (3).xlsx`** (abas `level 4` a `level 8`) "
      "para a hierarquia, e **`Draft 23 Out 26.pptx`** para os cargos do topo e "
      "para as duas caixas azuis na linha do CEO.")
    A("")
    A("**Leitura da planilha:** em cada aba, a pessoa da linha está no level da aba "
      "e a coluna B (`Nível de gerente N-1`) é o gestor direto dela. As colunas vêm "
      "em células mescladas, então o nome vale para todo o bloco abaixo.")
    A("")
    A("**Correção aplicada:** o draft mostra duas caixas azuis ligadas ao CEO "
      "(`Prod | Aéreo | Sampaio` e `Prod | Marítimo | Filipe`) e risca com **X** as "
      "caixas dessas duas pessoas no organograma do COO. Então Produto Aéreo e "
      "Produto Marítimo **reportam ao CEO**, e não ao COO como está na planilha. "
      "Com isso o CEO administra **8 pessoas**, não 6. O passo a passo está em "
      "`METODOLOGIA.md`.")
    A("")

    # ---- topo
    A("## O topo da estrutura")
    A("")
    A("| Level | Nome | Cargo | Fonte do cargo | C-LEVEL no DRE | Diretos | Indiretos | Total |")
    A("|---|---|---|---|---|---|---|---|")
    for p in sorted([x for x in todos if x["level"] <= 3],
                    key=lambda x: (x["level"], -x["total"])):
        cargo, fonte = cargo_de(p)
        mv = " ⬆" if p["nome"] in MOVIDOS_PARA_CEO else ""
        A(f"| L{p['level']} | **{p['nome']}**{mv} | {cargo} | {fonte} "
          f"| {C_LEVEL_DRE.get(p['nome'], '—')} | {p['diretos']} | {p['indiretos']} "
          f"| **{p['total']}** |")
    A("")
    A("⬆ = subiu de level por causa das caixas azuis do draft. "
      "Na planilha Factorial os dois estão como L4 dentro de Operações.")
    A("")
    A("Dois cargos vêm truncados no print do Factorial: *Chief Information & Tec…* "
      "(Cleyton) e *Chief Strategy Finance …* (Ana Carolina). No DRE são **CITO** e "
      "**CFSO**. O CEO é **Co-**Chief Executive Officer; a folha de setembro tem "
      "outras duas pessoas em `OWNERSHIP` fora do organograma "
      "(Fernando Esteves Arruda, EXECUTIVE BOARD, e Eduardo Reis de Almeida).")
    A("")
    A("---")
    A("")

    # ---- indice
    A("## Quantos gestores em cada level")
    A("")
    A("| Level | Pessoas no level | Gestores | Administram juntos | Só executam |")
    A("|---|---|---|---|---|")
    for n in sorted(por_level):
        gr = por_level[n]
        g = [p for p in gr if p["diretos"] > 0]
        soma = sum(p["diretos"] for p in g)
        A(f"| **L{n}** | {len(gr)} | **{len(g)}** | "
          f"{f'{soma} pessoas do L{n+1}' if soma else '—'} | {len(gr) - len(g)} |")
    A(f"| **Total** | **{len(todos)}** | "
      f"**{sum(1 for p in todos if p['diretos'] > 0)}** | "
      f"{sum(p['diretos'] for p in todos)} | "
      f"{sum(1 for p in todos if p['diretos'] == 0)} |")
    A("")
    A("No último level ninguém administra ninguém — é o fim da cadeia.")
    A("")
    A("---")
    A("")

    # ---- blocos por level
    for n in sorted(por_level):
        gestores = [p for p in por_level[n] if p["diretos"] > 0]
        simples = [p for p in por_level[n] if p["diretos"] == 0]
        A(f"# LEVEL {n}")
        A("")
        if not gestores:
            A(f"**Nenhum gestor neste level.** As {len(simples)} pessoas do L{n} "
              "não administram ninguém.")
            A("")
            A(f"| Nome | Cargo | Gestor direto (L{n-1}) |")
            A("|---|---|---|")
            for p in sorted(simples, key=lambda x: x["nome"]):
                A(f"| {p['nome']} | {cargo_de(p)[0]} | {p['gestor']} |")
            A("")
            continue

        A(f"**{plural(len(gestores), 'gestor', 'gestores')}** no L{n}, administrando "
          f"{plural(sum(p['diretos'] for p in gestores), 'pessoa')} diretamente.")
        if simples:
            A("")
            A(f"(as outras {len(simples)} pessoas do L{n} não administram ninguém — "
              "listadas no fim do bloco)")
        A("")
        A("| Gestor | Cargo | Diretos | Indiretos | Total | Confere com o draft? |")
        A("|---|---|---|---|---|---|")
        for p in gestores:
            risco = " ✗" if p["nome"] in RISCADOS else ""
            A(f"| {p['nome']}{risco} | {cargo_de(p)[0]} | **{p['diretos']}** "
              f"| {p['indiretos']} | {p['total']} | {confere(p, como_esta)} |")
        if any(p["nome"] in RISCADOS for p in gestores):
            A("")
            A("✗ = caixa riscada com X no draft.")
        A("")

        for p in gestores:
            cargo, fonte = cargo_de(p)
            A(f"### {p['nome']}  ·  L{n}")
            A("")
            linha = cargo + (f"  *(cargo lido no {fonte})*" if fonte.startswith("draft")
                             else "")
            A(linha + "  ")
            A(f"Administra **{plural(p['diretos'], 'pessoa')}** diretamente"
              + (f" e {plural(p['indiretos'], 'pessoa')} indiretamente, "
                 f"somando {p['total']}." if p["indiretos"] else
                 " — nenhum indireto: o time todo é de executores.")
              + (f"  \nReporta a: **{p['gestor']}** (L{n-1})" if p["gestor"] else
                 "  \nTopo da estrutura — não reporta a ninguém."))
            if p["nome"] in MOVIDOS_PARA_CEO:
                A(f"  \n⬆ Reposicionado pelo draft: caixa azul "
                  f"`{MOVIDOS_PARA_CEO[p['nome']]}`. Na planilha Factorial "
                  f"reporta a {p['gestor_factorial']}.")
            if p["nome"] in RISCADOS:
                sl, pos, obs = RISCADOS[p["nome"]]
                A(f"  \n✗ Caixa riscada com X no slide {sl} ({pos}): {obs}.")
            A("")
            A("| # | Administra | Cargo | Level | Esse também administra? |")
            A("|---|---|---|---|---|")
            for i, f in enumerate(filhos[p["nome"]], 1):
                adm = (f"sim — {plural(f['diretos'], 'pessoa')}" if f["diretos"]
                       else "não")
                A(f"| {i} | {f['nome']} | {cargo_de(f)[0]} | L{f['level']} | {adm} |")
            A("")

        if simples:
            A(f"### L{n} — não administram ninguém ({len(simples)} pessoas)")
            A("")
            A(f"| Nome | Cargo | Gestor direto (L{n-1}) |")
            A("|---|---|---|")
            for p in sorted(simples, key=lambda x: x["nome"]):
                A(f"| {p['nome']} | {cargo_de(p)[0]} | {p['gestor']} |")
            A("")
        A("---")
        A("")

    # ---- divergencias resumidas
    div = dados["div"]
    gaps = [r for r in div if "gap real" in r["causa"]]
    A("# Divergências com as bases financeiras")
    A("")
    A("Detalhe linha por linha na aba **`Divergencias`** da planilha. Resumo:")
    A("")
    A("| Comparação | Casaram | Só no organograma | Só na base |")
    A("|---|---|---|---|")
    for b in ("Baseline DRE 26.2", "Folha Set/26"):
        c = sum(1 for r in div if r["base"] == b and r["tipo"] == "casou")
        o = sum(1 for r in div if r["base"] == b and r["tipo"] == "só no organograma")
        s = sum(1 for r in div if r["base"] == b and r["tipo"].startswith("só n")
                and r["tipo"] != "só no organograma")
        A(f"| Organograma × {b} | {c if c else '—'} | {o} | {s} |")
    A("")
    A("Quase toda diferença é data de corte: o organograma é de outubro e inclui "
      "admissões até 16/nov; a folha é de setembro; o DRE é de agosto. "
      "Descontado isso, mais quem está em aviso prévio, os sócios, as linhas de "
      "consultoria e a linha duplicada, sobram as pessoas abaixo — "
      "**com custo lançado e sem posição no organograma**:")
    A("")
    agrupado = {}
    for r in gaps:
        k = " ".join(toks(r["nome_base"]))
        agrupado.setdefault(k, []).append(r)
    A("| Nome | Como aparece na folha de set/26 | Como aparece no DRE |")
    A("|---|---|---|")
    for k, rs in agrupado.items():
        f = next((r for r in rs if r["base"] == "Folha Set/26"), None)
        d = next((r for r in rs if r["base"] == "Baseline DRE 26.2"), None)
        nome = (d or f)["nome_base"].title()
        A(f"| **{nome}** "
          f"| {f['info1'] + ' · ' + f['info2'] if f else 'não está'} "
          f"| {d['info1'] + ' · ' + d['info2'] + ' · ' + d['info3'] if d else 'não está'} |")
    A("")
    A(f"São {len(agrupado)} pessoas: duas de desenvolvimento de indústrias e um "
      "estagiário de Ocean Operations. Todas têm custo lançado nas duas bases "
      "e nenhuma tem gestor no Factorial.")
    A("")
    A(f"E `{TESTE}`, um registro de teste com admissão em 2022 cadastrado como "
      "*Air Freight Design Specialist* sob Rafael Araujo Belfiore: infla o "
      "organograma em 1 pessoa.")
    A("")
    A("---")
    A("")
    A("## Anotações encontradas no draft")
    A("")
    A("| Slide(s) | O que está anotado em cima do print |")
    A("|---|---|")
    for sl, nota in NOTAS_DRAFT:
        A(f"| {sl} | {str(nota).replace('|', '/')} |")
    A("")
    A("---")
    A("")
    A("## Como regerar")
    A("")
    A("```bash")
    A("pip install openpyxl")
    A("python3 estrutura/gestores_por_level.py")
    A("```")
    return "\n".join(L)


# ------------------------------------------------------------------ excel
def gerar_xlsx(todos, filhos, ceo, como_esta, dados):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    por_level = defaultdict(list)
    for p in todos:
        por_level[p["level"]].append(p)
    area = diretoria_de(todos, filhos, ceo)

    # ---- 1. quem reporta a quem
    ws = wb.create_sheet("Quem reporta a quem")
    ws.append(["Level", "Nome", "Cargo", "Fonte do cargo", "Diretos", "Indiretos",
               "Total abaixo", "Gestor direto", "Level do gestor", "Diretoria (L3)",
               "Level na planilha Factorial", "Gestor na planilha Factorial",
               "Observação do draft", "Local", "Início"])
    for p in sorted(todos, key=lambda x: (x["level"], -x["diretos"], x["nome"])):
        cargo, fonte = cargo_de(p)
        obs = p["movido"]
        if p["nome"] in RISCADOS:
            sl, pos, o = RISCADOS[p["nome"]]
            obs = (obs + " | " if obs else "") + f"riscado com X no slide {sl}: {o}"
        ws.append([f"L{p['level']}", p["nome"], cargo, fonte, p["diretos"],
                   p["indiretos"], p["total"], p["gestor"] or "—",
                   f"L{p['level']-1}" if p["gestor"] else "—",
                   area.get(p["nome"], "—" if p["level"] <= 2 else ""),
                   f"L{p['level_factorial']}", p["gestor_factorial"] or "—",
                   obs, p["local"], p["inicio"]])

    # ---- 2. resumo por level
    ws = wb.create_sheet("Resumo por level")
    ws.append(["Level", "Pessoas", "Gestores", "Só executam",
               "Diretos administrados", "Maior time (diretos)"])
    for n in sorted(por_level):
        gr = por_level[n]
        g = [p for p in gr if p["diretos"] > 0]
        ws.append([f"L{n}", len(gr), len(g), len(gr) - len(g),
                   sum(p["diretos"] for p in g),
                   max([p["diretos"] for p in g], default=0)])
    ws.append(["Total", len(todos), sum(1 for p in todos if p["diretos"] > 0),
               sum(1 for p in todos if p["diretos"] == 0),
               sum(p["diretos"] for p in todos), ""])
    ws.append([])
    ws.append(["Diretoria (L3)", "Responsável", "Cargo", "Diretos", "Pessoas na área"])
    for l3 in filhos[ceo]:
        ws.append([l3["nome"], l3["nome"], cargo_de(l3)[0], l3["diretos"], l3["total"]])

    # ---- 3. uma aba por level
    for n in sorted(por_level):
        gestores = [p for p in por_level[n] if p["diretos"] > 0]
        if not gestores:
            continue
        ws = wb.create_sheet(f"L{n} gestores")
        ws.append([f"Gestor (L{n})", "Cargo do gestor", "Diretos", "Indiretos",
                   "Total abaixo", "Confere com o draft?", "Administra",
                   "Cargo de quem ele administra", "Level",
                   "Esse administra quantos?"])
        for p in gestores:
            for f in filhos[p["nome"]]:
                ws.append([p["nome"], cargo_de(p)[0], p["diretos"], p["indiretos"],
                           p["total"], confere(p, como_esta), f["nome"],
                           cargo_de(f)[0], f"L{f['level']}", f["diretos"]])

    # ---- 4. NOVA aba: divergencias
    ws = wb.create_sheet("Divergencias")
    ws.append(["Base comparada", "Situação", "Level", "Nome no organograma",
               "Cargo no organograma", "Gestor", "Nome na base", "C-LEVEL / Posição",
               "Área / Centro de custo", "Posição padrão / Situação",
               "Folha / Tipo de custo", "TT Custo/Mês", "Causa"])
    for r in dados["div"]:
        ws.append([r["base"], r["tipo"], r["level"], r["nome_org"], r["cargo_org"],
                   r["gestor"], r["nome_base"], r["info1"], r["info2"], r["info3"],
                   r["info4"], r["custo"], r["causa"]])

    # ---- formatacao
    cor = {2: "0D1F31", 3: "1B3A5C", 4: "2E6CA4", 5: "67A3D9", 6: "AED2F0",
           7: "DFEAF5", 8: "F4F8FC"}
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.fill = PatternFill("solid", fgColor="1B3A5C")
            c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}1"
        larg = {}
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None:
                    larg[c.column] = min(46, max(larg.get(c.column, 10),
                                                 len(str(c.value)) + 2))
        for col, w in larg.items():
            sheet.column_dimensions[get_column_letter(col)].width = w

    ws = wb["Quem reporta a quem"]
    for row in ws.iter_rows(min_row=2):
        n = int(str(row[0].value)[1:])
        row[0].fill = PatternFill("solid", fgColor=cor.get(n, "F4F8FC"))
        row[0].font = Font(bold=True, color="FFFFFF" if n <= 4 else "06243D")
        row[0].alignment = Alignment(horizontal="center")
    ws = wb["Divergencias"]
    tons = {"só na folha": "FDE2E2", "só no DRE": "FDF0DC",
            "só no organograma": "E3F0FD", "casou": "FFFFFF"}
    for row in ws.iter_rows(min_row=2):
        f = PatternFill("solid", fgColor=tons.get(str(row[1].value), "FFFFFF"))
        for c in row[:2]:
            c.fill = f
        if "gap real" in str(row[12].value) or "teste" in str(row[12].value):
            row[12].font = Font(bold=True, color="9C0006")

    destino = os.path.join(SAIDA, "Gestores_por_Level.xlsx")
    wb.save(destino)
    return destino


def main():
    fac = ler_factorial()
    # arvore da planilha pura, usada para conferir os numeros impressos no draft
    t0, f0, ceo0 = montar(fac, aplicar_draft=False)
    como_esta = {p["nome"]: p["total"] for p in t0}
    # arvore com as caixas azuis do draft aplicadas
    todos, filhos, ceo = montar(fac, aplicar_draft=True)

    dados = {
        "base": ler_aba("Base_Setembro26 (1).xlsx", "Setembro Unificado", "NOME"),
        "dre": ler_aba("Baseline DRE 26.2.xlsx", "Baseline ativos agosto", "Colaborador"),
    }
    dados["div"] = divergencias(todos, dados["base"], dados["dre"])

    cam = os.path.join(SAIDA, "Gestores_por_Level.md")
    with open(cam, "w", encoding="utf-8") as fh:
        fh.write(gerar_md(todos, filhos, ceo, como_esta, dados) + "\n")
    xl = gerar_xlsx(todos, filhos, ceo, como_esta, dados)

    print(f"{len(todos)} pessoas | CEO administra {filhos[ceo].__len__()} diretos")
    for n in sorted({p['level'] for p in todos}):
        g = [p for p in todos if p["level"] == n and p["diretos"] > 0]
        print(f"  L{n}: {sum(1 for p in todos if p['level']==n):>3} pessoas, "
              f"{len(g):>2} gestores")
    conf = confere_draft(t0)
    print(f"draft: {len(conf)} gestores conferidos, "
          f"{sum(1 for c in conf if c['dif'])} com diferença")
    print(f"divergencias: {len(dados['div'])} linhas")
    print("ok ->", cam)
    print("ok ->", xl)


if __name__ == "__main__":
    main()
