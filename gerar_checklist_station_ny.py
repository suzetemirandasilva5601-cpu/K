# -*- coding: utf-8 -*-
"""Checklist acionavel do modelo de station-piloto NY, por dimensao. Excel replicavel."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = 'DMS - Modelo Station Piloto NY - Checklist.xlsx'
AZUL = 'FF005E9E'
AZULCLARO = 'FFE5F0F8'
VERDE = 'FFE8F5E9'
LARANJA = 'FFFDF0E4'
CINZA = 'FFF2F2F2'
AMAR = 'FFFFF2CC'
branco = Font(color='FFFFFFFF', bold=True, size=11)
bold = Font(bold=True)
tit = Font(bold=True, size=13, color=AZUL)
ital = Font(italic=True, size=8, color='FF555555')
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin = Side(style='thin', color='FFBFBFBF')
bd = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()

# ---- ABA 1: operacao minima (MVP) ----
s0 = wb.active
s0.title = 'Operacao Minima'
s0['A1'] = 'Operacao minima para a base operar (gate de go-live)'
s0['A1'].font = tit
s0.merge_cells('A1:D1')
s0['A2'] = 'Abaixo destes 7 elementos nao ha go-live. Comunicacao, benchmark e healthcare entram depois.'
s0['A2'].font = ital
s0.merge_cells('A2:D2')
hdr = ['#', 'Elemento minimo', 'Por que e piso', 'Status']
for j, h in enumerate(hdr, 1):
    c = s0.cell(row=4, column=j, value=h); c.font = branco
    c.fill = PatternFill('solid', fgColor=AZUL); c.alignment = center; c.border = bd
mvp = [
    (1, 'Parceiro CCSF homologado + security TSA aprovado', 'sem screening a carga nao embarca'),
    (2, 'Capacidade confirmada p/ ao menos 1 industria (aerospace)', 'base sem carga atendivel nao opera'),
    (3, 'Analista local contratado + espaco fisico no parceiro', 'sem operacao local nao replica Houston'),
    (4, 'Cadastro do parceiro no sistema de pagamento', 'sem isso nao paga nem fatura'),
    (5, 'SOP default da praca clonada de Houston', 'sem SOP cada embarque e improviso'),
    (6, 'SLA de informacao com o parceiro definido', 'AOG sem timing falha no 1o teste'),
    (7, 'Carga-piloto identificada (cliente + shipper reais)', 'go-live se prova com carga real'),
]
for i, (n, e, p) in enumerate(mvp, 5):
    for j, v in enumerate((n, e, p, 'Nao iniciado'), 1):
        c = s0.cell(row=i, column=j, value=v); c.border = bd
        c.alignment = center if j in (1, 4) else left
        c.font = Font(size=9)
        if j == 4:
            c.fill = PatternFill('solid', fgColor=AMAR)
for col, w in zip('ABCD', (5, 48, 45, 16)):
    s0.column_dimensions[col].width = w

# ---- ABA 2: checklist por dimensao ----
s1 = wb.create_sheet('Checklist Pre e Pos')
s1['A1'] = 'Checklist de abertura de station por dimensao (pre -> go-live -> pos)'
s1['A1'].font = tit
s1.merge_cells('A1:E1')
s1['A2'] = 'Replicavel: clonar por praca. Cada item tem dono e status. Fonte: reuniao NYC 27/10 + modelo Houston.'
s1['A2'].font = ital
s1.merge_cells('A2:E2')
cols = ['Fase', 'Dimensao', 'Item', 'Dono sugerido', 'Status']
for j, h in enumerate(cols, 1):
    c = s1.cell(row=4, column=j, value=h); c.font = branco
    c.fill = PatternFill('solid', fgColor=AZUL); c.alignment = center; c.border = bd

itens = [
    ('Pre', 'Fornecedores', 'Prospeccao de CCSF neutro na regiao do aeroporto', 'RD'),
    ('Pre', 'Fornecedores', 'Homologacao de seguranca (CPET) aprovada', 'RD + parceiro'),
    ('Pre', 'Fornecedores', 'Security / TSA screening validado', 'RD + parceiro'),
    ('Pre', 'Fornecedores', 'Capacidade por industria confirmada (nao so declarada)', 'RD + Operacao'),
    ('Pre', 'Fornecedores', 'Contrato/SLA assinado e cadastrado no sistema (TS)', 'RD + Financeiro'),
    ('Pre', 'Fornecedores', 'Plano B de parceiro identificado', 'RD'),
    ('Pre', 'Contabil/Financeiro', 'Cadastro do parceiro no pagamento (buying/selling)', 'Financeiro'),
    ('Pre', 'Contabil/Financeiro', 'Contabilizacao do pagamento ao parceiro no exterior', 'Financeiro'),
    ('Pre', 'Contabil/Financeiro', 'Spread cambial e prepaid abroad tratados', 'Financeiro'),
    ('Pre', 'Contabil/Financeiro', 'Previsao de servicos e faturamento (template de custo)', 'Faturamento'),
    ('Pre', 'Contabil/Financeiro', 'Custo de overtime/fim de semana na conta de AOG', 'Financeiro'),
    ('Pre', 'Regulatorio/Compliance', 'TSA / screening aprovado', 'Compliance + RD'),
    ('Pre', 'Regulatorio/Compliance', 'Requisitos de exportacao por industria mapeados', 'Compliance'),
    ('Pre', 'Regulatorio/Compliance', 'Validacao de compliance (alinhar com Patricia)', 'Compliance'),
    ('Pre', 'Regulatorio/Compliance', 'Restricoes por tipo de carga (DG, refrigerado, perecivel)', 'Compliance + Operacao'),
    ('Pre', 'Operacao', 'Visita operacional (estrutura fisica, pesagem, distancia)', 'Operacao'),
    ('Pre', 'Operacao', 'Espaco in-house para o analista', 'Operacao'),
    ('Pre', 'Operacao', 'Processos de fim de semana / overtime acordados', 'Operacao'),
    ('Pre', 'Operacao', 'SLA de informacao (pre-alerta, WR, CCT, pesagem, AOG)', 'Operacao'),
    ('Pre', 'Operacao', 'SOP default da praca clonada de Houston', 'Operacao/CAN'),
    ('Pre', 'Operacao', 'Reuniao de operacao com todos os times', 'Operacao'),
    ('Pre', 'Operacao', 'Mapeamento de transportadoras (benchmark)', 'Operacao'),
    ('Pre', 'Pessoas (RH)', 'Contratacao do analista local', 'Operacao + RH'),
    ('Pre', 'Comunicacao corporativa', 'Definir dono (TMT ou Operacao)', 'TMT'),
    ('Pre', 'Comunicacao corporativa', 'Anuncio da operacao ao mercado', 'TMT'),
    ('Pre', 'Comunicacao corporativa', 'Material comercial para os KAMs', 'TMT + KAM'),
    ('Pre', 'Carga/Comercial', 'Mapeamento de clientes e shippers que migram', 'Produto + RD'),
    ('Pre', 'Carga/Comercial', 'Carga-piloto identificada (cliente+shipper reais)', 'Produto'),
    ('Pre', 'Carga/Comercial', 'Volume e faturamento previstos da 1a onda', 'Produto + Faturamento'),
    ('Go-live', 'Operacao', '1o trial com carga real controlada (AOG Bristol via JFK)', 'Operacao'),
    ('Go-live', 'Operacao', 'Fluxo ponta a ponta validado', 'Operacao'),
    ('Go-live', 'Sistema', 'Sistema reconhecendo a origem nova (operacionalizacao)', 'Operacao + TI'),
    ('Pos', 'Acompanhamento', '% de AOG no prazo', 'Operacao + Qualidade'),
    ('Pos', 'Acompanhamento', 'Volume migrado vs previsto', 'Operacao'),
    ('Pos', 'Acompanhamento', 'Qualidade documental (erros/retrabalho)', 'Qualidade'),
    ('Pos', 'Acompanhamento', 'Margem realizada vs prevista', 'Financeiro'),
    ('Pos', 'Acompanhamento', 'Custo por embarque JFK vs Miami', 'Financeiro'),
    ('Pos', 'Consolidacao', 'Registro do tempo real de cada etapa', 'ATY/Operacao'),
    ('Pos', 'Consolidacao', 'Itens que travaram e por que (prevencao)', 'ATY/Operacao'),
    ('Pos', 'Consolidacao', 'Separar o especifico de NY vs clonavel', 'ATY'),
    ('Pos', 'Consolidacao', 'Decisao sobre healthcare (2o parceiro?)', 'RD + Operacao'),
    ('Pos', 'Estrategico', 'Impacto sobre Miami (desidratacao?)', 'Operacao'),
    ('Pos', 'Estrategico', 'A base se pagou no horizonte previsto?', 'Financeiro'),
]
fase_cor = {'Pre': AZULCLARO, 'Go-live': LARANJA, 'Pos': VERDE}
r = 5
for fase, dim, item, dono in itens:
    for j, v in enumerate((fase, dim, item, dono, 'Nao iniciado'), 1):
        c = s1.cell(row=r, column=j, value=v); c.border = bd
        c.alignment = left if j in (2, 3, 4) else center
        c.font = Font(size=9)
        if j == 1:
            c.fill = PatternFill('solid', fgColor=fase_cor[fase]); c.font = bold
        if j == 5:
            c.fill = PatternFill('solid', fgColor=AMAR)
    r += 1
dv = DataValidation(type='list', formula1='"Nao iniciado,Em andamento,Concluido,Bloqueado,N/A"', allow_blank=True)
s1.add_data_validation(dv); dv.add('E5:E%d' % (r - 1))
for col, w in zip('ABCDE', (10, 24, 52, 24, 15)):
    s1.column_dimensions[col].width = w
s1.freeze_panes = 'A5'

wb.save(SAIDA)
print('Excel gerado:', SAIDA, '| itens:', len(itens))
