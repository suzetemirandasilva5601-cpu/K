#!/usr/bin/env python3
"""
Dados lidos do Draft 23 Out 26.pptx e rotinas de conferencia contra as bases
financeiras. Modulo de apoio para gestores_por_level.py.

Fontes:
  Colaboradores-factorial (3).xlsx  -> organograma, abas 'level 4' a 'level 8'
  Draft 23 Out 26.pptx              -> 65 prints do Factorial + anotacoes manuais
  Base_Setembro26 (1).xlsx          -> aba 'Setembro Unificado' (folha de set/26)
  Baseline DRE 26.2.xlsx            -> aba 'Baseline ativos agosto' (custo/pessoa)
  Vagas 21-10.xlsx                  -> aba 'Vagas'
"""
import os
import re
import unicodedata
from collections import Counter, defaultdict

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NIVEL_ABA = {"level 4": 4, "level 5": 5, "level 6": 6, "level 7": 7, "level 8": 8}
SEM_CARGO = "(cargo não informado na planilha)"
CORTE_FOLHA = "2026-09-30"      # ultimo dia da competencia Set/26
DATA_DRAFT = "2026-10-23"
TESTE = "Colaborador Teste"
PAT_INTERN = re.compile(r"\b(intern|apprentice)\b", re.I)


# ============================================================ DADOS DO DRAFT
# Cargos lidos nos organogramas. Os levels 2 e 3 nao tem linha propria na
# planilha Factorial, so aparecem como nome de gestor: estes vieram do slide 2.
CARGOS_DRAFT = {
    "Jose Luiz de Oliveira Belixior Junior": ("Co-Chief Executive Officer", "s2"),
    "Raioni De Oliveira Santos": ("Chief Operating Officer", "s2/s46"),
    "Andre Lemos Tavares": ("Chief Commercial Officer", "s2/s17"),
    "Cleyton Cardoso Esteves": ("Chief Information & Tec…", "s2/s24"),
    "Ana Carolina Botelho Caldeira Vasconcellos": ("Chief Strategy Finance …", "s2/s29"),
    "Carolina Karnakis de Nobrega": ("Head of Talent", "s2/s37"),
    "Luis Fernando Santos Muniz": ("Head of Services", "s2/s43"),
}

# Sigla de C-LEVEL usada no DRE e nas Vagas.
C_LEVEL_DRE = {
    "Jose Luiz de Oliveira Belixior Junior": "Co-CEO",
    "Raioni De Oliveira Santos": "COO",
    "Andre Lemos Tavares": "CCO",
    "Cleyton Cardoso Esteves": "CITO",
    "Ana Carolina Botelho Caldeira Vasconcellos": "CFSO",
    "Carolina Karnakis de Nobrega": "dentro de Co-CEO (Talent Management Team)",
    "Luis Fernando Santos Muniz": "dentro de Co-CEO (Services)",
    # os dois Produtos continuam dentro de COO no DRE, que e a visao da planilha
    "Filipe Teixeira Sampaio": "COO / Air Product (ainda sob o COO no DRE)",
    "Filipe Carneiro De Oliveira": "COO / Ocean Product (ainda sob o COO no DRE)",
}

# Caixas AZUIS desenhadas na linha do CEO (slides 2, 15, 22, 27, 35, 41, 44),
# ligadas ao CEO por conector manual. Movem o Produto Aereo e o Maritimo de
# baixo do COO para reporte direto ao CEO.
MOVIDOS_PARA_CEO = {
    "Filipe Teixeira Sampaio": "Prod | Aéreo | Sampaio",
    "Filipe Carneiro De Oliveira": "Prod | Marítimo | Filipe",
}

# Caixas RISCADAS COM X no draft. Posicao conferida mapeando as coordenadas
# das linhas do PowerPoint sobre os pixels da imagem e recortando a regiao.
RISCADOS = {
    "Filipe Carneiro De Oliveira": ("46", "x≈307px",
                                    "sai de baixo do COO — vira caixa azul no CEO"),
    "Filipe Teixeira Sampaio": ("46", "x≈1525px",
                                "sai de baixo do COO — vira caixa azul no CEO"),
    "Magno Bezerra Rocha": ("48", "x≈501px",
                            "riscado sob Carolina Bento Mendoza, sem caixa de destino"),
    "Gabriella De Melo Pereira": ("54", "x≈466px",
                                  "riscada sob Daniela Bueno, sem caixa de destino"),
    "Pamela Iris Gregorio Ibarra": ("54", "x≈770px",
                                    "riscada sob Daniela Bueno, sem caixa de destino"),
}

# 'liderados' impressos em cada organograma, com o slide onde aparecem.
DRAFT_LIDERADOS = {
    "Jose Luiz de Oliveira Belixior Junior": (330, "s2"),
    "Raioni De Oliveira Santos": (199, "s2/s46"),
    "Cleyton Cardoso Esteves": (47, "s2/s24"),
    "Ana Carolina Botelho Caldeira Vasconcellos": (33, "s2/s29"),
    "Carolina Karnakis de Nobrega": (23, "s2/s37"),
    "Andre Lemos Tavares": (17, "s2/s17"),
    "Luis Fernando Santos Muniz": (5, "s2/s43"),
    "Carolina Bento Mendoza": (82, "s46/s48"),
    "Filipe Teixeira Sampaio": (28, "s46/s4"),
    "Natalia Correia Da Costa Parreira": (35, "s46/s58"),
    "Filipe Carneiro De Oliveira": (18, "s46/s11"),
    "Marco Antonio Maraia Villa": (12, "s46/s71"),
    "Possidonio Dantas de Almeida Neto": (8, "s46/s65"),
    "Jonathas Amancio Alves": (6, "s46/s75"),
    "Lucy Maia De Moraes": (2, "s46"),
    "Victor Telles Gonzaga": (37, "s24/s26"),
    "Julio Cesar Francisco de Souza": (8, "s24/s25"),
    "Roberta Duarte Correia": (18, "s29/s30"),
    "Daniel de Carvalho Herrera Braga": (8, "s29/s34"),
    "Rodrigo Fontes Dominciano": (8, "s37/s39"),
    "Vivian Ramalho De Assumpçao": (5, "s37/s38"),
    "Leandro Teciano": (3, "s37/s40"),
    "Cristiane Martins Soares Giovanelli": (1, "s37/s40"),
    "Katharine Ferreira De Souza": (1, "s39"),
    "Ivan Garcia Castro": (2, "s17/s18"),
    "Yuri Da Silva Torreao": (3, "s17/s19"),
    "Ricardo Tebexreni Jakowatz": (5, "s17/s20"),
    "Jose Carlos Duarte Freitas Sobrinho": (2, "s17/s21"),
    "Nathalia Santos Reis Cordeiro": (2, "s19"),
    "Danielle Moreira Rabelo Machado": (1, "s20"),
    "Marcelo Portantiolo de Lima": (2, "s20"),
    "Joao Pedro Da Costa Marcicano Rodrigues": (23, "s4/s5"),
    "Tayna Neves": (3, "s4/s9"),
    "Dandara Cristina Santos Ru": (7, "s5/s6"),
    "Rafael Araujo Belfiore": (7, "s5/s7"),
    "Vinicius Teixeira Studart Pereira": (5, "s5/s8"),
    "Rony Ghiu De Moraes Rego Da Silva": (3, "s11/s12"),
    "Ivy Kaori Oishi": (6, "s11/s13"),
    "Luiz Antonio Duarte Ferreira Junior": (5, "s11/s14"),
    "Daniela Bueno dos Reis Craig": (47, "s48/s54"),
    "Thais Pinho Cosenza": (27, "s48/s50"),
    "Julia Sena Campelo Lopes": (4, "s48/s49"),
    "Rubia Furtado": (12, "s50/s53"),
    "Leonardo Bonamichi Vaz de Lima": (8, "s50/s52"),
    "Isabelle Carolina De Melo Martinez Lizama": (3, "s50/s51"),
    "Pamela Iris Gregorio Ibarra": (21, "s54/s69"),
    "Gabriella De Melo Pereira": (16, "s54/s67"),
    "Adriane Borel Cordeiro Impagliazzo": (5, "s54/s55"),
    "Charlete Alfaia Guimaraes Santana": (1, "s54/s56"),
    "Thais Espindola Oliveira Mathias": (8, "s69"),
    "Julia Olegario Varella": (1, "s69"),
    "Julia Heinrichs Pimentel": (6, "s67/s68"),
    "Lenice Maria da Costa Barbosa": (1, "s67"),
    "Kayza Cristina Silva Pereira": (12, "s58/s60"),
    "Leonardo Araujo de Menezes": (7, "s58/s63"),
    "Ricardo Dos Reis Castro": (4, "s58/s61"),
    "Augusto Cesar Sotelo Medeiros": (4, "s58/s62"),
    "Andrea Oliveira Brito Do Couto": (2, "s58/s59"),
    "Jesus Guillermo Ledi Mendoza": (5, "s65/s66"),
    "Caio Vitor de Carvalho Santos": (4, "s71/s73"),
    "Natalie Giseli De Maria Correa": (2, "s71/s72"),
    "Camila Evelyn De Souza Lima": (8, "s30/s32"),
    "Gisele Ferreira Pinto": (5, "s30/s33"),
    "Andressa Franco Da Silva": (2, "s30/s31"),
}

# Anotacoes de texto e agrupamentos encontrados em cima dos prints.
NOTAS_DRAFT = [
    ("2, 15, 22, 27, 35, 41, 44",
     "Duas caixas azuis na linha do CEO: 'Prod / Aéreo / Sampaio' e "
     "'Prod / Marítimo / Filipe', ligadas ao CEO por conector manual"),
    ("46", "X riscando Filipe Carneiro De Oliveira e Filipe Teixeira Sampaio "
           "no organograma do COO"),
    ("48", "X riscando Magno Bezerra Rocha (Airport Operations Manager)"),
    ("54", "X riscando Gabriella De Melo Pereira e Pamela Iris Gregorio Ibarra"),
    ("6, 14", "RD — 'A estrutura é separada por região: EUA / EUROPA / ÁSIA. "
              "Precisamos fazer essa abertura'"),
    ("7, 8, 13", "FD — 'A estrutura é separada por Cliente. Precisamos fazer "
                 "essa abertura'"),
    ("20", "Três retângulos agrupando as frentes Vertical / Auto / High Tech, "
           "e uma caixa 'Vaga de TL' na frente de High Tech"),
    ("75", "Dois retângulos agrupando o time de Jonathas Amancio Alves em "
           "'KAM' (5 Key Account Executives) e 'Capacity' (1 pessoa)"),
    ("24", "'A estrutura foi ajustada recentemente pela Cleyton — temos que "
           "pegar a visão atualizada'"),
    ("29", "'A estrutura foi ajustada recentemente pela Ana Carol — temos que "
           "pegar a visão atualizada'"),
    ("37", "'A estrutura foi ajustada recentemente pela Karnakis — temos que "
           "pegar a visão atualizada'"),
    ("46 ", "'A estrutura foi ajustada recentemente pelo Raioni — temos que "
            "pegar a visão atualizada'"),
]


# ============================================================ NOMES
STOP = {"DE", "DA", "DO", "DOS", "DAS", "E", "JUNIOR", "JR", "NETO", "FILHO",
        "SOBRINHO", "CONSULTOR"}


def toks(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z ]", " ", s).upper()
    return [t for t in s.split() if t not in STOP and len(t) > 1]


def score(a, b):
    """Pontua a chance de dois nomes serem a mesma pessoa. 0 = descartado."""
    ta, tb = toks(a), toks(b)
    if not ta or not tb:
        return 0
    A, B = set(ta), set(tb)
    comum = A & B
    if not comum:
        return 0
    if ta[0] != tb[0] and ta[-1] != tb[-1]:
        return 0
    if len(comum) < 2 and not (A <= B or B <= A):
        return 0
    s = 10 * len(comum)
    s += 6 if ta[0] == tb[0] else 0
    s += 4 if ta[-1] == tb[-1] else 0
    s += 3 if (A <= B or B <= A) else 0
    return s - abs(len(A) - len(B))


def emparelha(la, ca, lb, cb):
    """Casa duas listas pelo melhor score global, nao pela primeira ocorrencia."""
    cand = sorted(((score(x[ca], y[cb]), i, j)
                   for i, x in enumerate(la) for j, y in enumerate(lb)),
                  key=lambda t: -t[0])
    ua, ub, pares = set(), set(), []
    for s, i, j in cand:
        if s <= 0 or i in ua or j in ub:
            continue
        ua.add(i)
        ub.add(j)
        pares.append((la[i], lb[j], s))
    return (pares,
            [x for i, x in enumerate(la) if i not in ua],
            [y for j, y in enumerate(lb) if j not in ub])


# ============================================================ LEITURA
def ler_factorial():
    wb = openpyxl.load_workbook(
        os.path.join(RAIZ, "Colaboradores-factorial (3).xlsx"), data_only=True)
    out = []
    for aba, nivel in NIVEL_ABA.items():
        ws = wb[aba]
        g_acima = g_direto = None
        for r in ws.iter_rows(min_row=3, values_only=True):
            if all(c is None for c in r):
                continue
            if r[0]:
                g_acima = str(r[0]).strip()
            if r[1]:
                g_direto = str(r[1]).strip()
            out.append({
                "nome": f"{r[2] or ''} {r[3] or ''}".strip(),
                "level_factorial": nivel,
                "cargo": (str(r[4]).strip() if r[4] else ""),
                "gestor_factorial": g_direto, "gestor_acima": g_acima,
                "diretos_planilha": int(r[6] or 0), "total_planilha": int(r[5] or 0),
                "genero": r[8] or "", "local": r[9] or "",
                "inicio": str(r[10])[:10] if r[10] else "",
            })
    return out


def ler_aba(arquivo, aba, col_nome):
    wb = openpyxl.load_workbook(os.path.join(RAIZ, arquivo), data_only=True)
    ws = wb[aba]
    hdr = [c.value for c in ws[1]]
    i = hdr.index(col_nome)
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)
            if r[i] not in (None, "")]


# ============================================================ ARVORE
def montar(fac, aplicar_draft=True):
    """Monta a arvore. Com aplicar_draft, move Produto Aereo e Maritimo para o CEO
    (as duas caixas azuis do draft) e recalcula o level como profundidade."""
    nomes = {p["nome"] for p in fac}
    l3 = sorted({p["gestor_factorial"] for p in fac if p["level_factorial"] == 4} - nomes)
    ceo = sorted({p["gestor_acima"] for p in fac if p["level_factorial"] == 4} - nomes)[0]

    todos = [dict(p) for p in fac]
    for nome in l3:
        todos.append({"nome": nome, "level_factorial": 3, "cargo": "",
                      "gestor_factorial": ceo, "gestor_acima": None,
                      "diretos_planilha": 0, "total_planilha": 0,
                      "genero": "", "local": "", "inicio": ""})
    todos.append({"nome": ceo, "level_factorial": 2, "cargo": "",
                  "gestor_factorial": None, "gestor_acima": None,
                  "diretos_planilha": 0, "total_planilha": 0,
                  "genero": "", "local": "", "inicio": ""})

    for p in todos:
        p["gestor"] = p["gestor_factorial"]
        p["movido"] = ""
        if aplicar_draft and p["nome"] in MOVIDOS_PARA_CEO:
            p["gestor"] = ceo
            p["movido"] = (f"caixa azul do draft: saiu de {p['gestor_factorial']} "
                           f"para o CEO")

    filhos = defaultdict(list)
    for p in todos:
        if p["gestor"]:
            filhos[p["gestor"]].append(p)

    # level = profundidade a partir do CEO (CEO = 2, como na planilha)
    idx = {p["nome"]: p for p in todos}
    fila = [(ceo, 2)]
    while fila:
        nome, lv = fila.pop()
        idx[nome]["level"] = lv
        for f in filhos[nome]:
            fila.append((f["nome"], lv + 1))

    def total(n):
        t = len(filhos[n])
        for f in filhos[n]:
            t += total(f["nome"])
        return t

    for p in todos:
        p["diretos"] = len(filhos[p["nome"]])
        p["total"] = total(p["nome"])
        p["indiretos"] = p["total"] - p["diretos"]
    for k in filhos:
        filhos[k].sort(key=lambda x: (-x["diretos"], x["nome"]))
    return todos, filhos, ceo


def diretoria_de(todos, filhos, ceo):
    """Mapeia cada pessoa para a area de L3 em que ela esta."""
    out = {}
    for l3 in filhos[ceo]:
        pilha = [l3]
        while pilha:
            x = pilha.pop()
            out[x["nome"]] = l3["nome"]
            pilha.extend(filhos[x["nome"]])
    return out


# ============================================================ DIVERGENCIAS
def classifica_organograma(x):
    if x["nome"] == TESTE:
        return "Registro de teste — excluir da base"
    if x["level"] == 2:
        return "Co-CEO: não entra na folha nem no DRE"
    if (x["inicio"] or "9999") > CORTE_FOLHA:
        return f"Admitido em {x['inicio']}, depois do corte da base"
    return "SEM CUSTO NA BASE — verificar"


def indice_folha(base):
    """Agrupa a folha por nome normalizado, guardando quantas linhas cada um tem."""
    idx = defaultdict(list)
    for b in base:
        idx[" ".join(toks(b["NOME"]))].append(b)
    return idx


def classifica_folha(y, idx):
    s = str(y["SITUAÇÃO"])
    if len(idx[" ".join(toks(y["NOME"]))]) > 1:
        return ("Linha duplicada na folha (dois contratos para a mesma pessoa) — "
                "ela está no organograma")
    if s in ("AVISO PREVIO V", "OFF BOARDING"):
        return f"{s} — saindo, já retirada do organograma"
    if y["TIPO DE CUSTO"] == "Sócios":
        return "Sócio / Executive Board — fora do organograma"
    if y["CONTRATO"] == "CONSULTOR":
        return "Contrato de consultoria lançado como pessoa"
    if s == "TRANSIÇÃO DE CONTRATO":
        return "Transição de contrato — linha duplicada"
    return "ATIVO SEM POSIÇÃO NO ORGANOGRAMA — gap real"


def classifica_dre(y, idx):
    """Classifica quem esta no DRE e nao no organograma, olhando a folha."""
    nome = str(y["Colaborador"])
    if str(y["ARÉA (Centro de Custo)"]).endswith("Consultoria") or nome.isupper():
        return "Linha de consultoria, não é pessoa do organograma"
    if nome in CARGOS_DRAFT:
        return "L2/L3 — não tem linha própria na planilha Factorial"
    na_folha = idx.get(" ".join(toks(nome)))
    if not na_folha:
        return "Estava no baseline de agosto e saiu antes de setembro"
    s = str(na_folha[0]["SITUAÇÃO"])
    if s in ("AVISO PREVIO V", "OFF BOARDING"):
        return f"{s} na folha de setembro — saindo, já retirada do organograma"
    if na_folha[0]["TIPO DE CUSTO"] == "Sócios":
        return "Sócio / Executive Board — fora do organograma"
    return "ATIVO SEM POSIÇÃO NO ORGANOGRAMA — gap real"


def divergencias(todos, base, dre):
    """Lista unica de divergencias contra o DRE e contra a folha de setembro."""
    linhas = []
    idx = indice_folha(base)

    pd_, soF_dre, soD = emparelha(todos, "nome", dre, "Colaborador")
    for f, d, s in pd_:
        linhas.append({
            "base": "Baseline DRE 26.2", "tipo": "casou", "level": f"L{f['level']}",
            "nome_org": f["nome"], "cargo_org": f["cargo"] or SEM_CARGO,
            "gestor": f["gestor"] or "—", "nome_base": d["Colaborador"],
            "info1": d["C-LEVEL"], "info2": d["ARÉA (Centro de Custo)"],
            "info3": d["POSIÇÃO PADRÃO 26.2"], "info4": d["Folha/Colchão"],
            "custo": d["TT Custo/Mês"], "causa": "",
        })
    for f in soF_dre:
        linhas.append({
            "base": "Baseline DRE 26.2", "tipo": "só no organograma",
            "level": f"L{f['level']}", "nome_org": f["nome"],
            "cargo_org": f["cargo"] or SEM_CARGO, "gestor": f["gestor"] or "—",
            "nome_base": "", "info1": "", "info2": "", "info3": "",
            "info4": f"início {f['inicio'] or '—'}", "custo": "",
            "causa": classifica_organograma(f),
        })
    for d in soD:
        linhas.append({
            "base": "Baseline DRE 26.2", "tipo": "só no DRE", "level": "",
            "nome_org": "", "cargo_org": "", "gestor": "",
            "nome_base": d["Colaborador"], "info1": d["C-LEVEL"],
            "info2": d["ARÉA (Centro de Custo)"], "info3": d["POSIÇÃO PADRÃO 26.2"],
            "info4": d["Folha/Colchão"], "custo": d["TT Custo/Mês"],
            "causa": classifica_dre(d, idx),
        })

    pb, soF_b, soB = emparelha(todos, "nome", base, "NOME")
    for f in soF_b:
        linhas.append({
            "base": "Folha Set/26", "tipo": "só no organograma",
            "level": f"L{f['level']}", "nome_org": f["nome"],
            "cargo_org": f["cargo"] or SEM_CARGO, "gestor": f["gestor"] or "—",
            "nome_base": "", "info1": "", "info2": "", "info3": "",
            "info4": f"início {f['inicio'] or '—'}", "custo": "",
            "causa": classifica_organograma(f),
        })
    for b in soB:
        linhas.append({
            "base": "Folha Set/26", "tipo": "só na folha", "level": "",
            "nome_org": "", "cargo_org": "", "gestor": "", "nome_base": b["NOME"],
            "info1": b["POSIÇÃO"], "info2": b["AREA"], "info3": b["SITUAÇÃO"],
            "info4": b["TIPO DE CUSTO"], "custo": "",
            "causa": classifica_folha(b, idx),
        })
    ordem = {"só na folha": 0, "só no DRE": 1, "só no organograma": 2, "casou": 3}
    linhas.sort(key=lambda r: (ordem[r["tipo"]], r["base"], r["nome_org"] or r["nome_base"]))
    return linhas


def confere_draft(todos):
    """Compara os liderados impressos no draft com o calculado na planilha."""
    out = []
    for p in todos:
        if p["nome"] not in DRAFT_LIDERADOS:
            continue
        d, slide = DRAFT_LIDERADOS[p["nome"]]
        out.append({"nome": p["nome"], "level": p["level"], "draft": d,
                    "slide": slide, "planilha": p["total"], "dif": p["total"] - d})
    return sorted(out, key=lambda r: (r["level"], -r["planilha"]))
