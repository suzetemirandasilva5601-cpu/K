# -*- coding: utf-8 -*-
"""
Template replicavel de abertura de station (SOP), em Excel.
Baseado na reuniao Operacao NYC (27/10) e no Checklist NYC existente,
mas generalizado para clonar em qualquer localidade.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SAIDA = 'DMS - Template Abertura de Base (SOP replicavel).xlsx'
AZUL = 'FF005E9E'
AZULCLARO = 'FFE5F0F8'
CINZA = 'FFF2F2F2'
AMARELO = 'FFFFF2CC'
branco = Font(color='FFFFFFFF', bold=True, size=11)
bold = Font(bold=True)
tit = Font(bold=True, size=13, color=AZUL)
ital = Font(italic=True, size=8, color='FF555555')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin = Side(style='thin', color='FFBFBFBF')
bd = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
sh = wb.active
sh.title = 'Abertura de Base'

sh['A1'] = 'DMS - Template de Abertura de Base / Station (SOP replicavel)'
sh['A1'].font = tit
sh.merge_cells('A1:G1')
sh['A2'] = ('Preencher Localidade/Parceiro/Datas no topo e usar como checklist. '
            'Clonar esta aba para cada nova praca. Primeira aplicacao: Nova York (JFK).')
sh['A2'].font = ital
sh.merge_cells('A2:G2')

# cabecalho de contexto (editavel)
ctx = [('Localidade / Aeroporto', 'Nova York / JFK'),
       ('Parceiro (CCSF/warehouse)', 'a definir (neutro, proximo a JFK)'),
       ('Modelo', 'Parceria de warehousing + analista local DMS (espelho Houston)'),
       ('Cargas-alvo 1a onda', 'Aerospace (AOG) · Healthcare condicionado a refrigeracao'),
       ('Data-alvo de go-live', 'a definir (depende do TSA)')]
r = 4
for k, v in ctx:
    sh.cell(row=r, column=1, value=k).font = bold
    sh.cell(row=r, column=1).alignment = left
    c = sh.cell(row=r, column=2, value=v)
    c.fill = PatternFill('solid', fgColor=AMARELO)
    c.alignment = left
    sh.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    r += 1

# tabela de acoes
hr = r + 1
cols = ['Fase', 'Acao', 'Descricao / criterio de pronto', 'Dono sugerido', 'Depende de', 'Prazo', 'Status']
for j, c in enumerate(cols, 1):
    cell = sh.cell(row=hr, column=j, value=c)
    cell.font = branco
    cell.fill = PatternFill('solid', fgColor=AZUL)
    cell.alignment = center
    cell.border = bd

linhas = [
    ('Pre-Operacao', 'Prospeccao de parceiro', 'Achar CCSF/warehouse neutro na regiao do aeroporto', 'RD', '-', '', 'Concluido'),
    ('Pre-Operacao', 'Homologacao de seguranca', 'Formulario de security/CPET aprovado (TSA/screening)', 'RD + parceiro', 'Prospeccao', '', 'Em andamento'),
    ('Pre-Operacao', 'Qualificacao do parceiro', 'Capacidade por industria (aerospace OK; checar refrigeracao p/ healthcare)', 'RD + Operacao', 'Prospeccao', '', ''),
    ('Pre-Operacao', 'Mapeamento de cargas e clientes', 'Que cargas migram (regiao+tipo), shippers, clientes, volume e faturamento previsto', 'RD + Produto', '-', '', ''),
    ('Pre-Operacao', 'Visita operacional', 'Validar estrutura fisica: espaco analista, recebimento, pesagem, distancia aeroporto, overtime', 'Operacao', 'Prospeccao', '', ''),
    ('Pre-Operacao', 'Contratacao do analista local', 'Headcount no modelo Houston/Roberta', 'Operacao + RH', '-', '', ''),
    ('Pre-Operacao', 'Cadastro no sistema de pagamento', 'Parceiro cadastrado (buying/selling) - DEFINIR quem faz', 'Financeiro', 'Qualificacao', '', ''),
    ('Pre-Operacao', 'Espaco in-house para o analista', 'Posicao fisica dentro do parceiro', 'Operacao', 'Visita', '', ''),
    ('Pre-Operacao', 'SOP default da station', 'Clonar SOP default de Houston; depois CAN clona por shipper/cliente', 'Operacao/CAN + Qualidade', '-', '', ''),
    ('Pre-Operacao', 'Aspectos regulatorios/fiscais', 'TSA, screening, requisitos de exportacao por industria; contabil do pagamento no exterior', 'Financeiro + Compliance', '-', '', ''),
    ('Pre-Operacao', 'Previsao de servicos e faturamento', 'Template de custo (funcionario, armazem, servico RD) - hoje ad hoc', 'Faturamento', 'Mapeamento', '', ''),
    ('Pre-Operacao', 'SLA de informacao com o parceiro', 'Timings de pre-alerta, warehouse receipt, CCT, pesagem, AOG', 'Operacao', 'Visita', '', ''),
    ('Pre-Operacao', 'Reuniao de operacao', 'Alinhar CCT, pesagem, dinamica de AOG com todos os times', 'Operacao', 'SLA', '', ''),
    ('Pre-Operacao', 'Comunicacao corporativa', 'Anunciar nova operacao ao mercado (radar de embarque)', 'TMT', 'Go-live', '', ''),
    ('Piloto / Go-live', 'Primeiro trial controlado', 'Carga real (ex.: AOG Bristol via JFK), nao cotacao; fluxo ponta a ponta', 'Operacao', 'TSA + SOP', '', ''),
    ('Pos-Operacao', 'Acompanhamento de performance', 'SLA, % AOG no prazo, qualidade documental, volume migrado vs previsto, margem', 'Operacao + Qualidade', 'Go-live', '', ''),
]
rr = hr + 1
fase_cor = {'Pre-Operacao': AZULCLARO, 'Piloto / Go-live': 'FFFDF0E4', 'Pos-Operacao': 'FFE8F5E9'}
for ln in linhas:
    for j, v in enumerate(ln, 1):
        c = sh.cell(row=rr, column=j, value=v)
        c.border = bd
        c.alignment = left if j in (2, 3) else center
        c.font = Font(size=9)
        if j == 1:
            c.fill = PatternFill('solid', fgColor=fase_cor.get(v, CINZA))
            c.font = bold
    rr += 1

# dropdown de status
dv = DataValidation(type='list', formula1='"Nao iniciado,Em andamento,Concluido,Bloqueado"', allow_blank=True)
sh.add_data_validation(dv)
dv.add('G%d:G%d' % (hr + 1, rr - 1))

for col, w in zip('ABCDEFG', (15, 30, 50, 22, 16, 12, 14)):
    sh.column_dimensions[col].width = w
sh.freeze_panes = 'A%d' % (hr + 1)

# aba 2: os 6 blocos do template generico
s2 = wb.create_sheet('Template 6 blocos')
s2['A1'] = 'Os 6 blocos do template replicavel (o que clona em qualquer praca)'
s2['A1'].font = tit
s2.merge_cells('A1:C1')
blocos = [
    ('1. Comunicacao corporativa', 'Quando e como anunciar a station ao mercado', 'TMT'),
    ('2. Cadastro do fornecedor', 'Onde e quem cadastra o parceiro no pagamento (buying/selling)', 'Financeiro'),
    ('3. Aspectos contabeis', 'Pagamento ao parceiro no exterior; spread cambial; prepaid abroad', 'Financeiro'),
    ('4. Aspectos regulatorios', 'TSA, screening (CCSF), requisitos de exportacao por industria', 'Compliance + RD'),
    ('5. Pre-operacao', 'Checklist A1-A14 com donos e interdependencias (Gantt)', 'Operacao'),
    ('6. Pos-operacao', 'Rotina de performance e revisao de SLA', 'Operacao + Qualidade'),
]
for j, h in enumerate(['Bloco', 'O que cobre', 'Dono'], 1):
    c = s2.cell(row=3, column=j, value=h)
    c.font = branco
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.alignment = center
    c.border = bd
for i, (a, b, c) in enumerate(blocos, 4):
    for j, v in enumerate((a, b, c), 1):
        cell = s2.cell(row=i, column=j, value=v)
        cell.border = bd
        cell.alignment = left
        if j == 1:
            cell.font = bold
for col, w in zip('ABC', (30, 60, 22)):
    s2.column_dimensions[col].width = w

wb.save(SAIDA)
print('Excel gerado:', SAIDA)
