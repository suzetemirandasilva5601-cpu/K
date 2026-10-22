# -*- coding: utf-8 -*-
"""Gera Word (análise) e Excel (conciliação) da diferença DRE agosto x Folha setembro - DMS."""
import unicodedata, re
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

DRE_FILE = "Baseline DRE 26.2.xlsx"
DRE_SHEET = "Baseline ativos agosto"
FOL_FILE = "Base_Setembro26 (1).xlsx"
FOL_SHEET = "Setembro Unificado"

def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFKD", str(s).strip()).encode("ascii", "ignore").decode("ascii")
    s = re.sub(r"[^A-Za-z ]", " ", s)
    return re.sub(r"\s+", " ", s).upper().strip()

def first_last(s):
    t = norm(s).split()
    if not t: return ""
    return t[0] + "|" + t[-1] if len(t) > 1 else t[0]

def load(f, sheet, namecol):
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb[sheet]
    rows = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=namecol).value
        if v in (None, ""): continue
        rows.append({"linha": r, "nome": str(v).strip(), "norm": norm(v), "fl": first_last(v)})
    return rows

dre = load(DRE_FILE, DRE_SHEET, 3)
fol = load(FOL_FILE, FOL_SHEET, 5)

n_dre, n_fol = len(dre), len(fol)
dif = n_dre - n_fol

# duplicatas na folha
dup_fol = {k: v for k, v in Counter(x["norm"] for x in fol).items() if v > 1}

# match por primeiro+ultimo nome
fol_fl = set(x["fl"] for x in fol)
dre_fl = set(x["fl"] for x in dre)

saidas = [x for x in dre if x["fl"] not in fol_fl]      # DRE e nao folha
entradas = [x for x in fol if x["fl"] not in dre_fl]    # folha e nao DRE

# grafia divergente conhecida (conferir manualmente) - pares provaveis
grafia = [
    ("PIRACI", "CONSULTOR - PIRACI OLIVEIRA SOCIEDADE DE ADVOGADOS",
     "Mesma PJ/pessoa com razao social na folha"),
    ("Isabella FRA", "ISABELLA STODIEK",
     "Nome truncado na DRE; confirmar se e a mesma pessoa"),
    ("AUAD ENGENHARIA DE SOFTWARE LTDA", "(verificar pessoa fisica na folha)",
     "Razao social de PJ na DRE; pode estar na folha pelo nome da pessoa"),
]

# ============================ EXCEL ============================
wb = openpyxl.Workbook()

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

def autofit(ws):
    for col in ws.columns:
        width = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                width = max(width, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(width + 3, 12), 60)

# --- Aba Resumo ---
ws = wb.active
ws.title = "Resumo"
ws["A1"] = "Conciliacao DRE Agosto x Folha Setembro - DMS"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A1:C1")
resumo = [
    ("Indicador", "Valor", "Observacao"),
    ("DRE de agosto (colaboradores)", n_dre, "Linhas 2-329 da aba 'Baseline ativos agosto'"),
    ("Folha de setembro (colaboradores)", n_fol, "Linhas 2-324 da aba 'Setembro Unificado'"),
    ("Diferenca liquida (DRE - Folha)", dif, "NAO sao 3 linhas; sao 5"),
    ("Nomes so na DRE (sairam/nao considerados)", len(saidas), "Rotatividade - saidas"),
    ("Nomes so na Folha (entradas)", len(entradas), "Rotatividade - entradas"),
    ("Duplicata na folha", "; ".join(dup_fol.keys()) or "nenhuma",
     "PAMELA IRIS GREGORIO IBARRA aparece 2x" if dup_fol else ""),
    ("Reconciliacao", f"{n_dre} - {len(saidas)} + {len(entradas)} = {n_dre - len(saidas) + len(entradas)}",
     "Aproxima de 322/323 (duplicata na folha)"),
]
start = 3
for i, (a, b, c) in enumerate(resumo):
    ws.cell(row=start + i, column=1, value=a)
    ws.cell(row=start + i, column=2, value=b)
    ws.cell(row=start + i, column=3, value=c)
style_header(ws, start, 3)
for r in range(start + 1, start + len(resumo)):
    for c in range(1, 4):
        ws.cell(row=r, column=c).border = BORDER
autofit(ws)

# --- Aba Sairam ---
ws2 = wb.create_sheet("Sairam (so na DRE)")
ws2.append(["#", "Nome na DRE (agosto)", "Linha na DRE"])
style_header(ws2, 1, 3)
for i, x in enumerate(sorted(saidas, key=lambda z: z["nome"].upper()), 1):
    ws2.append([i, x["nome"], x["linha"]])
    for c in range(1, 4): ws2.cell(row=i + 1, column=c).border = BORDER
autofit(ws2)

# --- Aba Entraram ---
ws3 = wb.create_sheet("Entraram (so na Folha)")
ws3.append(["#", "Nome na Folha (setembro)", "Linha na Folha"])
style_header(ws3, 1, 3)
for i, x in enumerate(sorted(entradas, key=lambda z: z["nome"].upper()), 1):
    ws3.append([i, x["nome"], x["linha"]])
    for c in range(1, 4): ws3.cell(row=i + 1, column=c).border = BORDER
autofit(ws3)

# --- Aba Grafia divergente ---
ws4 = wb.create_sheet("Grafia a conferir")
ws4.append(["#", "Como aparece na DRE", "Provavel na Folha", "Observacao"])
style_header(ws4, 1, 4)
for i, (a, b, obs) in enumerate(grafia, 1):
    ws4.append([i, a, b, obs])
    for c in range(1, 5): ws4.cell(row=i + 1, column=c).border = BORDER
autofit(ws4)

xlsx_out = "Conciliacao_DRE_Agosto_x_Folha_Setembro.xlsx"
wb.save(xlsx_out)
print("Excel gerado:", xlsx_out)

# ============================ WORD ============================
doc = Document()
# estilo base
style = doc.styles["Normal"]
style.font.name = "Calibri"; style.font.size = Pt(11)

t = doc.add_heading("Analise da Diferenca - DRE de Agosto x Folha de Setembro", level=0)

p = doc.add_paragraph()
p.add_run("Empresa: ").bold = True; p.add_run("DMS")
p = doc.add_paragraph()
p.add_run("Objeto: ").bold = True
p.add_run("Correlacao entre a DRE (baseline de agosto) e a folha de pagamento de setembro.")

doc.add_heading("1. Numeros encontrados", level=1)
tbl = doc.add_table(rows=1, cols=2); tbl.style = "Light Grid Accent 1"
hdr = tbl.rows[0].cells
hdr[0].text = "Base"; hdr[1].text = "Colaboradores"
for base, qtd in [("DRE de agosto (linhas 2-329)", str(n_dre)),
                  ("Folha de setembro (linhas 2-324)", str(n_fol)),
                  ("Diferenca liquida (DRE - Folha)", str(dif))]:
    row = tbl.add_row().cells
    row[0].text = base; row[1].text = qtd
    if base.startswith("Diferenca"):
        for cell in row:
            for par in cell.paragraphs:
                for run in par.runs: run.bold = True

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Conclusao do item: ").bold = True
p.add_run("a diferenca real entre as duas bases e de ")
r = p.add_run("5 linhas (328 - 323), e nao 3."); r.bold = True

doc.add_heading("2. O que existe de verdade: rotatividade nos dois sentidos", level=1)
doc.add_paragraph(
    "A diferenca liquida de apenas 5 linhas esconde um movimento bem maior de pessoas. "
    "Entre agosto e setembro houve rotatividade nos dois sentidos - gente que saiu e gente "
    "que entrou - e os dois efeitos quase se cancelam na contagem total.")
doc.add_paragraph(
    "~23 nomes aparecem na DRE de agosto e NAO na folha de setembro "
    "(saidas / colaboradores nao considerados no mes).", style="List Bullet")
doc.add_paragraph(
    "~16-17 nomes aparecem na folha de setembro e NAO na DRE de agosto (entradas).",
    style="List Bullet")

doc.add_heading("3. Como a conta fecha", level=1)
p = doc.add_paragraph()
p.add_run("328 - 23 (sairam) + 17 (entraram) = 322").bold = True
p.add_run(", resultado proximo de 322/323.")
doc.add_paragraph(
    "A pequena divergencia de 1 se explica por um nome duplicado na folha de setembro: "
    "PAMELA IRIS GREGORIO IBARRA aparece 2 vezes.")

doc.add_heading("4. Pontos de atencao (grafia divergente)", level=1)
doc.add_paragraph(
    "Alguns nomes contados como 'so na DRE' na verdade estao na folha, apenas escritos de "
    "forma diferente. Devem ser conferidos manualmente antes de tratar como saida real:")
for a, b, obs in grafia:
    doc.add_paragraph(f"{a}  ->  {b}  ({obs})", style="List Bullet")

doc.add_heading("5. Resumo executivo", level=1)
doc.add_paragraph(
    "A alteracao entre a DRE de agosto e a folha de setembro NAO e de 3 linhas. Trata-se de "
    "uma diferenca liquida de 5 linhas, resultante de aproximadamente 23 saidas e 17 entradas "
    "- ou seja, rotatividade relevante no quadro, mascarada pela proximidade dos totais.")

docx_out = "Analise_Diferenca_DRE_Agosto_x_Folha_Setembro.docx"
doc.save(docx_out)
print("Word gerado:", docx_out)
print("\nResumo: DRE", n_dre, "| Folha", n_fol, "| dif", dif,
      "| saidas", len(saidas), "| entradas", len(entradas), "| dup", dup_fol)
