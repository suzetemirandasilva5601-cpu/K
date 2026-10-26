# -*- coding: utf-8 -*-
"""
Gera o Excel da arvore de Operacoes com as colunas:
  Agosto (realizado, DRE baseline) | Setembro (realizado, ativos) |
  Variacao ago->set | Vagas em aberto | Leitura

IMPORTANTE: a DRE de agosto e o REALIZADO de agosto, nao o orcado. Nao existe
quadro-alvo nos arquivos. As colunas estao rotuladas para deixar isso claro.
"""
import datetime
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SAIDA = 'DMS - Arvore Operacoes - Ago x Set x Vagas.xlsx'


def nok(v):
    return v is not None and str(v).strip() not in ('', 'None', '.', '-', 'TBD', 'TDB', '?')


REF = datetime.date(2026, 10, 21)


def est(r):
    if nok(r[6]):
        return 'preenchida'
    if isinstance(r[4], datetime.datetime) and r[4].date() <= REF:
        return 'a confirmar'
    return 'aberta'


def bl(c):
    c = str(c).upper()
    if 'WH ' in c or 'WAREHOUSE' in c:
        return 'EUA / Warehouses'
    if 'INLAND USA' in c or 'INLAND AMERICAS' in c:
        return 'EUA / Inland Freight America'
    if 'AMERICAS' in c:
        return 'EUA / Operacoes America'
    if 'INLAND EUROPA' in c:
        return 'Europa / Inland Freight Europa'
    if 'EUROPA' in c:
        return 'Europa / Operacoes Europa'
    if c.startswith('AIR CC') or c.startswith('OCEAN CC'):
        return 'Brasil / Customer Care'
    if 'EXPORT' in c:
        return 'Brasil / Exportacao'
    return None


# dados
dre = openpyxl.load_workbook('Baseline DRE 26.2.xlsx', data_only=True)['Baseline ativos agosto']
ago = Counter(bl(r[5]) for r in dre.iter_rows(min_row=2, values_only=True) if r[2] and bl(r[5]))
ws = openpyxl.load_workbook('Base_Setembro26 (1).xlsx', data_only=True)['Setembro Unificado']
A = [r for r in ws.iter_rows(min_row=2, values_only=True)
     if r[4] and str(r[6]).strip().upper() == 'ATIVO']
sett = Counter(bl(r[11]) for r in A if bl(r[11]))
wv = openpyxl.load_workbook('Vagas 21-10.xlsx', data_only=True)['Vagas']
V = [r for r in wv.iter_rows(min_row=2, values_only=True)
     if r[3] and str(r[3]).strip() not in ('', 'None')]
vag = Counter(bl(r[3]) for r in V if est(r) == 'aberta' and bl(r[3]))

GRUPOS = [
    ('EUA', ['EUA / Inland Freight America', 'EUA / Operacoes America', 'EUA / Warehouses']),
    ('Europa', ['Europa / Inland Freight Europa', 'Europa / Operacoes Europa']),
    ('Brasil', ['Brasil / Customer Care', 'Brasil / Exportacao']),
]

# ---- estilos ----
AZUL = 'FF005E9E'
AZULCLARO = 'FFE5F0F8'
CINZA = 'FFF2F2F2'
VERDE = 'FF2E7D32'
VERMELHO = 'FFC62828'
branco = Font(color='FFFFFFFF', bold=True, name='Calibri', size=11)
bold = Font(bold=True, name='Calibri')
normal = Font(name='Calibri')
center = Alignment(horizontal='center', vertical='center')
left = Alignment(horizontal='left', vertical='center')
thin = Side(style='thin', color='FFD0D7DE')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
sh = wb.active
sh.title = 'Arvore Operacoes'

# titulo
sh['A1'] = 'DMS - Arvore de Operacoes: Agosto x Setembro x Vagas em aberto'
sh['A1'].font = Font(bold=True, size=13, color=AZUL)
sh.merge_cells('A1:F1')
sh['A2'] = ('Ago e Set sao REALIZADOS (headcount ativo de cada mes). Nao existe quadro orcado '
            'nos arquivos; a variacao mostra a mudanca de quadro, nao orcado x realizado.')
sh['A2'].font = Font(italic=True, size=9, color='FF555555')
sh.merge_cells('A2:F2')

# cabecalho
hdr = ['Bloco / estrutura', 'Agosto (realizado)', 'Setembro (realizado)',
       'Variacao ago->set', 'Vagas em aberto', 'Leitura']
r0 = 4
for j, h in enumerate(hdr, 1):
    c = sh.cell(row=r0, column=j, value=h)
    c.font = branco
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.alignment = center if j > 1 else left
    c.border = border

row = r0 + 1
ta = ts = tv = 0
for nome_g, blocos in GRUPOS:
    # subtotal do grupo
    ga = sum(ago.get(b, 0) for b in blocos)
    gs = sum(sett.get(b, 0) for b in blocos)
    gv = sum(vag.get(b, 0) for b in blocos)
    ta += ga
    ts += gs
    tv += gv
    cells = [nome_g, ga, gs, gs - ga, gv, '']
    for j, val in enumerate(cells, 1):
        c = sh.cell(row=row, column=j, value=val)
        c.font = bold
        c.fill = PatternFill('solid', fgColor=AZULCLARO)
        c.alignment = left if j in (1, 6) else center
        c.border = border
        if j == 4 and isinstance(val, int):
            c.font = Font(bold=True, color=VERDE if val > 0 else (VERMELHO if val < 0 else '000000'))
    row += 1
    # blocos do grupo
    for b in blocos:
        a = ago.get(b, 0)
        s = sett.get(b, 0)
        v = vag.get(b, 0)
        var = s - a
        lt = ('cresceu %d' % var if var > 0 else ('caiu %d' % -var if var < 0 else 'estavel'))
        lt += ', %d vaga(s) em aberto' % v if v else ', sem vaga'
        vals = ['    ' + b.split(' / ')[1], a, s, var, v, lt]
        for j, val in enumerate(vals, 1):
            c = sh.cell(row=row, column=j, value=val)
            c.font = normal
            c.alignment = left if j in (1, 6) else center
            c.border = border
            if j == 4 and isinstance(val, int):
                c.font = Font(color=VERDE if val > 0 else (VERMELHO if val < 0 else '000000'))
        row += 1

# total
tot = ['TOTAL Operacoes', ta, ts, ts - ta, tv, '']
for j, val in enumerate(tot, 1):
    c = sh.cell(row=row, column=j, value=val)
    c.font = branco
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.alignment = left if j in (1, 6) else center
    c.border = border

# larguras
for j, w in enumerate([30, 18, 20, 18, 16, 34], 1):
    sh.column_dimensions[get_column_letter(j)].width = w

# nota de rodape
nr = row + 2
sh.cell(row=nr, column=1, value='Fonte: Baseline DRE 26.2 (ativos agosto); Base Setembro/26 (ativos); Vagas 21-10 (abertas).')
sh.cell(row=nr, column=1).font = Font(italic=True, size=8, color='FF555555')
sh.cell(row=nr + 1, column=1, value='Vaga em aberto = sem ocupante e sem admissao ja ocorrida. Variacao = Set - Ago (rotatividade liquida).')
sh.cell(row=nr + 1, column=1).font = Font(italic=True, size=8, color='FF555555')

wb.save(SAIDA)
print('Excel gerado:', SAIDA)
print('TOTAL: Ago %d | Set %d | Var %+d | Vagas %d' % (ta, ts, ts - ta, tv))
