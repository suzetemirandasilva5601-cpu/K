# -*- coding: utf-8 -*-
"""
Auditoria da apresentacao gerada.

Recalcula cada numero a partir das tres fontes e confere se aparece
corretamente no .pptx. Nao reaproveita nada do script gerador.
"""
import datetime
import re
import sys
from collections import Counter

import openpyxl
import pandas as pd
from pptx import Presentation

PPTX = 'DMS - Vagas em aberto x Ocupação nas Operações.pptx'
BASE = 'Base_Setembro26 (1).xlsx'
VAGAS = 'Vagas 21-10.xlsx'
CIS = 'CIS_20261013_ORDERS.parquet'
REF = datetime.date(2026, 10, 21)

falhas = []
checks = 0


def ck(desc, cond):
    global checks
    checks += 1
    if cond:
        print('  [OK]   %s' % desc)
    else:
        print('  [ERRO] %s' % desc)
        falhas.append(desc)


def num(v):
    if v is None or v == '-':
        return 0.0
    try:
        return float(v)
    except Exception:
        return 0.0


def cheio(v):
    return v is not None and str(v).strip() not in ('', 'None', '.', '-', 'TBD', 'TDB', '?')


def estado(r):
    if cheio(r[6]):
        return 'preenchida'
    a = r[4]
    if isinstance(a, datetime.datetime) and a.date() <= REF:
        return 'a confirmar'
    return 'aberta'


def bloco(c):
    c = str(c).upper()
    if 'WH ' in c or 'WAREHOUSE' in c:
        return 'WH'
    if 'INLAND USA' in c or 'INLAND AMERICAS' in c:
        return 'INLAND_US'
    if 'AMERICAS' in c:
        return 'OPS_AME'
    if 'INLAND EUROPA' in c:
        return 'INLAND_EU'
    if 'EUROPA' in c:
        return 'OPS_EU'
    if c.startswith('AIR CC') or c.startswith('OCEAN CC'):
        return 'CC'
    if 'EXPORT' in c:
        return 'EXP'
    return None


# ---------------- fontes ----------------
ws = openpyxl.load_workbook(BASE, data_only=True)['Setembro Unificado']
A = [r for r in ws.iter_rows(min_row=2, values_only=True)
     if r[4] and str(r[6]).strip().upper() == 'ATIVO']
wv = openpyxl.load_workbook(VAGAS, data_only=True)['Vagas']
T = [r for r in wv.iter_rows(min_row=2, values_only=True)
     if r[3] and str(r[3]).strip() not in ('', 'None')]
AB = [r for r in T if estado(r) == 'aberta']
PR = [r for r in T if estado(r) == 'preenchida']
CF = [r for r in T if estado(r) == 'a confirmar']
COO = [r for r in AB if str(r[1]).strip() == 'COO']

LIDER = {'TEAM LEADER', 'SR. TEAM LEADER', 'COORDINATOR', 'HEAD'}
b_at, b_tl, b_ab, b_semst = Counter(), Counter(), Counter(), Counter()
for r in A:
    b = bloco(r[11])
    if b:
        b_at[b] += 1
        if r[10] in LIDER:
            b_tl[b] += 1
for r in AB:
    b = bloco(r[3])
    if b:
        b_ab[b] += 1
        if 'ABERTA' not in str(r[7]).upper():
            b_semst[b] += 1

d = pd.read_parquet(CIS)
d['created_at'] = pd.to_datetime(d['created_at'])
p = d[(d.created_at >= '2026-04-01') & (d.created_at <= '2026-09-30')]
E = lambda f: p[f]['assembly_id'].nunique() / 6
EUA = round(E((p['off'].isin(['DMS MIAMI', 'DMS HOUSTON'])) & (p['mode'] == 'air')))
MIA_S = round(E((p['off'] == 'DMS MIAMI') & (p['mode'] == 'sea')))
EUR = round(E((p['off'].isin(['DMS FRANKFURT', 'DMS LOGISTICS LISBOA'])) & (p['mode'] == 'air')))
CC = round(E((p.imp_exp == 'IMP') & (p['mode'] == 'air')))
EX = round(E((p.imp_exp == 'EXP') & (p['mode'] == 'air')))
EXS = round(E((p.imp_exp == 'EXP') & (p['mode'] == 'sea')))
h = p[(p['off'] == 'DMS HOUSTON') & (p['mode'] == 'air')]
HJS = round(h[h.created_at >= '2026-07-01']['assembly_id'].nunique() / 3)
HAJ = round(h[h.created_at < '2026-07-01']['assembly_id'].nunique() / 3)

at_eua = b_at['INLAND_US'] + b_at['OPS_AME'] + b_at['WH']
at_eur = b_at['INLAND_EU'] + b_at['OPS_EU']
VPA = {'EUA': round((EUA + MIA_S) / at_eua), 'EUR': round(EUR / at_eur),
       'CC': round(CC / b_at['CC']), 'EXP': round((EX + EXS) / b_at['EXP'])}

TOT_AT = sum(b_at[k] for k in ['INLAND_US', 'OPS_AME', 'WH', 'INLAND_EU', 'OPS_EU', 'CC', 'EXP'])
TOT_AB = sum(b_ab[k] for k in ['INLAND_US', 'OPS_AME', 'WH', 'INLAND_EU', 'OPS_EU', 'CC', 'EXP'])

# ---------------- texto da apresentacao ----------------
prs = Presentation(PPTX)
txt_por_slide = {}
for i, s in enumerate(prs.slides, 1):
    pedacos = []
    for sh in s.shapes:
        if sh.has_text_frame:
            pedacos.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                pedacos.append(' | '.join(c.text for c in row.cells))
    txt_por_slide[i] = '\n'.join(pedacos)
TUDO = '\n'.join(txt_por_slide.values())


def tem(s, slide=None):
    alvo = TUDO if slide is None else txt_por_slide[slide]
    return s in alvo


def pcbr(v, t):
    return ('%.1f%%' % (v / t * 100)).replace('.', ',')


def brl(v):
    return 'R$ ' + '{:,.0f}'.format(v).replace(',', '.')


print('=' * 74)
print('AUDITORIA DA APRESENTACAO')
print('=' * 74)

print('\n[estrutura]')
ck('5 slides', len(prs.slides) == 5)
ck('slide 1 e capa (Blank_Analytics)', prs.slides[0].slide_layout.name == 'Blank_Analytics')
ck('slides 2 a 4 usam DMS_Conteudo',
   all(prs.slides[i].slide_layout.name == 'DMS_Conteudo' for i in (1, 2, 3)))
ck('slide 5 e encerramento sem texto',
   not txt_por_slide[5].strip())
ck('selo "Para discussao" nos 3 slides de conteudo',
   all('Para discussão' in txt_por_slide[i] for i in (2, 3, 4)))
ck('fonte citada nos 3 slides de conteudo',
   all('CIS ORDERS.parquet' in txt_por_slide[i] for i in (2, 3, 4)))

print('\n[vagas: classificacao]')
ck('49 vagas em aberto citado', tem('49') and len(AB) == 49)
ck('65 registros citado', tem('65 registros') and len(T) == 65)
ck('13 preenchidas citado', tem('13 já estão preenchidas') and len(PR) == 13)
ck('3 a confirmar citado', tem('3 dependem de confirmação') and len(CF) == 3)
ck('soma confere', len(AB) + len(PR) + len(CF) == len(T))

print('\n[COO]')
ck('35 vagas do COO', len(COO) == 35 and tem('35'))
ck('71,4%% do total', pcbr(len(COO), len(AB)) == '71,4%' and tem('71,4%'))
custo_coo = sum(num(r[12]) for r in COO)
custo_ab = sum(num(r[12]) for r in AB)
ck('custo COO %s' % brl(custo_coo), tem(brl(custo_coo)))
ck('%% custo COO %s' % pcbr(custo_coo, custo_ab), tem(pcbr(custo_coo, custo_ab)))
ck('custo total das abertas %s' % brl(custo_ab), tem(brl(custo_ab)))
aum = sum(1 for r in COO if 'Aumento' in str(r[15]))
ck('%d aumento de quadro no COO' % aum, tem('%d' % aum))
pub = sum(1 for r in COO if 'ABERTA' in str(r[7]).upper())
ck('%d de %d publicadas' % (pub, len(COO)), tem('%d de %d' % (pub, len(COO))))
area = Counter(str(r[2]).strip() for r in COO)
ck('Air Operations com %d vagas do COO' % area['Air Operations'],
   tem('Air Operations: %d' % area['Air Operations']))
ck('cada area do COO aparece',
   all(tem('%d' % v) for v in area.values()))

print('\n[blocos operacionais]')
ck('total de ativos %d' % TOT_AT, tem('%d' % TOT_AT, 3))
ck('total de vagas abertas %d' % TOT_AB, tem('%d' % TOT_AB, 3))
ck('%% total em aberto %s' % pcbr(TOT_AB, TOT_AT + TOT_AB), tem(pcbr(TOT_AB, TOT_AT + TOT_AB)))
rotulos = {'INLAND_EU': 'Inland Freight Europa', 'OPS_AME': 'Operações America',
           'INLAND_US': 'Inland Freight America', 'EXP': 'Exportação',
           'OPS_EU': 'Operações Europa', 'CC': 'Customer Care', 'WH': 'Warehouses'}
for k, rot in rotulos.items():
    den = b_at[k] + b_ab[k]
    pcv = pcbr(b_ab[k], den) if den else '-'
    ck('%s: %d ativos, %d vagas, %s' % (rot, b_at[k], b_ab[k], pcv),
       tem(pcv, 3) and tem(rot, 3))

print('\n[volumes do CIS]')
ck('embarques por ativo EUA %d' % VPA['EUA'], tem('%d' % VPA['EUA'], 3))
ck('embarques por ativo Europa %d' % VPA['EUR'], tem('%d' % VPA['EUR'], 3))
ck('embarques por ativo Customer Care %d' % VPA['CC'], tem('%d' % VPA['CC'], 3))
ck('embarques por ativo Exportacao %d' % VPA['EXP'], tem('%d' % VPA['EXP'], 3))
ck('Houston de %d para %d' % (HAJ, HJS), tem('de %d para %d' % (HAJ, HJS), 3))

print('\n[plano de acao]')
ck('7 frentes na tabela de acao', txt_por_slide[4].count('semanas') + txt_por_slide[4].count('Imediato') >= 7)
for k, rot in rotulos.items():
    ck('frente "%s" no plano' % rot, tem(rot, 4))
ck('%d de %d vagas de Operacoes sem estagio'
   % (sum(b_semst[k] for k in rotulos), TOT_AB),
   tem('%d das %d' % (sum(b_semst[k] for k in rotulos), TOT_AB), 4))

print('\n[qualidade do texto]')
EM, EN = '\u2014', '\u2013'
ck('nenhum travessao', EM not in TUDO and EN not in TUDO)
ck('percentuais com virgula decimal', not re.search(r'\d+\.\d+%', TUDO))
suspeitas = ['posicoes', 'estao', 'atencao', 'operacoes', 'ocupacao', 'lideranca',
             'confirmacao', 'contratacao', 'gestao', 'combinacao', 'reune']
ruins = [w for w in suspeitas if w in TUDO.lower()]
ck('acentuacao sem pendencias (%s)' % (ruins or 'nenhuma'), not ruins)

print('\n' + '=' * 74)
print('VERIFICACOES: %d' % checks)
if falhas:
    print('FALHAS: %d' % len(falhas))
    for f in falhas:
        print('  - %s' % f)
    sys.exit(1)
print('RESULTADO: todos os numeros da apresentacao conferem com as fontes.')
print('=' * 74)
