#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera "Apresentacoes para consultoria" na identidade visual do CCU.

O template e extraido do deck oficial "08. Financas para case 2026.2.pptx":
tema, layouts, fontes embutidas (Kanit / Kanit SemiBold / Work Sans), logo e
arte de capa sao reaproveitados do arquivo original.

Especificacao da identidade (medida no deck original):
  slide            10,00 x 5,625 in (16:9)
  vermelho CCU     #C1272D   (theme dk1)
  tinta            #161616   (theme lt1)
  cinza de apoio   #F3F3F3   (theme dk2)
  cinza de texto   #888888
  titulo de secao  Work Sans 24 pt vermelho em (0,395 ; 0,389)  9,21 x 0,63
  kicker           Work Sans 12 pt tinta    em (0,395 ; 1,080)
  rodape           Arial 9 pt               em y = 5,277
  cunha vermelha   rtTriangle 0,43 x 4,55 herdada do layout ONE_COLUMN_TEXT
  capa/encerram.   retangulo #161616 full bleed + borda branca de 2 pt

Icones: Material Symbols (Google, Apache License 2.0), rasterizados no
vermelho CCU e na tinta do tema. Ver apresentacao-consultoria/NOTICE.md.

Uso:
    python3 gerar_apresentacao_consultoria.py
"""

import os
import shutil
import zipfile
from copy import deepcopy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------- #
# caminhos
# --------------------------------------------------------------------------- #
BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(BASE, "08. Finanças para case 2026.2.pptx")
OUT = os.path.join(BASE, "CCU - Apresentacoes para consultoria.pptx")
PASTA = os.path.join(BASE, "apresentacao-consultoria")
ASSETS = os.path.join(PASTA, "assets")
ICONS = os.path.join(PASTA, "icons")

# --------------------------------------------------------------------------- #
# identidade visual CCU
# --------------------------------------------------------------------------- #
RED = RGBColor(0xC1, 0x27, 0x2D)
RED_DARK = RGBColor(0x91, 0x0F, 0x13)
INK = RGBColor(0x16, 0x16, 0x16)
GRAY = RGBColor(0x88, 0x88, 0x88)
LIGHT = RGBColor(0xF3, 0xF3, 0xF3)
BORDER = RGBColor(0xD9, 0xD9, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

F_DISP = "Kanit SemiBold"   # capa, divisores e agenda
F_NUM = "Kanit"             # numeros dos chips
F_BODY = "Work Sans"        # titulos de conteudo, kickers e corpo
F_FOOT = "Arial"            # rodape

DECK = "Apresentações para consultoria"
TURMA = "Prep4Consulting 2026.2 "
COPY = "Copyright © 2026 Clube de Consultoria Universitário. All rights reserved."

# geometria (polegadas)
TITLE_BOX = (0.3949, 0.3894, 9.2089, 0.6299)
KICKER_BOX = (0.3949, 1.0803, 9.2089, 0.3028)
FOOT_Y, FOOT_H = 5.2771, 0.3028
BODY_L, BODY_R = 0.52, 9.48           # margens livres da cunha vermelha
BODY_W = BODY_R - BODY_L

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# --------------------------------------------------------------------------- #
# ativos do deck original
# --------------------------------------------------------------------------- #
def extrair_assets():
    """Copia logo, arte de capa e icones de contato do deck oficial do CCU."""
    os.makedirs(ASSETS, exist_ok=True)
    mapa = {
        "image35.png": "logo-ccu.png",       # logo horizontal (fundo escuro)
        "image32.png": "arte-capa.png",      # arte da capa
        "image51.png": "icone-email.png",
        "image50.png": "icone-instagram.png",
        "image56.png": "icone-site.png",
        "image49.png": "icone-linkedin.png",
    }
    with zipfile.ZipFile(TEMPLATE) as z:
        for origem, destino in mapa.items():
            alvo = os.path.join(ASSETS, destino)
            if not os.path.exists(alvo):
                with z.open(f"ppt/media/{origem}") as f, open(alvo, "wb") as o:
                    shutil.copyfileobj(f, o)
    return {v: os.path.join(ASSETS, v) for v in mapa.values()}


def ic(nome, cor="red"):
    return os.path.join(ICONS, f"{nome}-{cor}.png")


# --------------------------------------------------------------------------- #
# utilitarios
# --------------------------------------------------------------------------- #
def limpar_slides(prs):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.get(qn("r:id")))
        lst.remove(sld)


def set_insets(tf, l=0.1, r=0.1, t=0.05, b=0.05):
    bp = tf._txBody.bodyPr
    bp.set("lIns", str(Emu(Inches(l)).emu))
    bp.set("rIns", str(Emu(Inches(r)).emu))
    bp.set("tIns", str(Emu(Inches(t)).emu))
    bp.set("bIns", str(Emu(Inches(b)).emu))


def limpar_tf(tf):
    p0 = tf.paragraphs[0]._p
    p0.getparent().remove(p0)


def par(tf, runs, size=12, color=INK, bold=False, font=F_BODY, align="l",
        line=1.0, before=0, after=0, italic=False):
    """runs: str ou lista de (texto, {size,color,bold,font,italic,link}).

    "\\n" dentro do texto vira quebra de linha real (<a:br/>), porque o
    PowerPoint ignora newline literal dentro de <a:t>.
    """
    p = tf.add_paragraph()
    p.alignment = ALIGN[align]
    p.line_spacing = line
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if isinstance(runs, str):
        runs = [(runs, {})]
    ultimo_rPr = None
    for txt, ov in runs:
        for k, trecho in enumerate(txt.split("\n")):
            if k > 0:
                br = p._p.makeelement(qn("a:br"), {})
                if ultimo_rPr is not None:
                    br.append(deepcopy(ultimo_rPr))
                p._p.append(br)
            r = p.add_run()
            r.text = trecho
            f = r.font
            f.size = Pt(ov.get("size", size))
            f.bold = ov.get("bold", bold)
            f.italic = ov.get("italic", italic)
            f.name = ov.get("font", font)
            if ov.get("link"):
                r.hyperlink.address = ov["link"]
            f.color.rgb = ov.get("color", color)
            f.underline = False
            ultimo_rPr = r._r.get_or_add_rPr()
    return p


def texto(slide, x, y, w, h, blocos, anchor="t", wrap=True,
          insets=(0.1, 0.1, 0.05, 0.05)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = ANCHOR[anchor]
    set_insets(tf, *insets)
    limpar_tf(tf)
    for b in blocos:
        par(tf, **b)
    return box


def forma(slide, tipo, x, y, w, h, fill=None, line=None, lw=1.0, adj=None):
    sh = slide.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    set_insets(tf, 0.06, 0.06, 0.03, 0.03)
    limpar_tf(tf)
    if adj:
        av = sh._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
        for nome, valor in adj:
            av.append(av.makeelement(qn("a:gd"),
                                     {"name": nome, "fmla": f"val {valor}"}))
    return sh


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0):
    return forma(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, lw)


def rect_txt(slide, x, y, w, h, runs, fill=None, line=None, lw=1.0,
             size=11, color=WHITE, bold=False, font=F_BODY, align="c",
             anchor="m", line_spc=1.05):
    sh = rect(slide, x, y, w, h, fill, line, lw)
    sh.text_frame.vertical_anchor = ANCHOR[anchor]
    par(sh.text_frame, runs, size=size, color=color, bold=bold, font=font,
        align=align, line=line_spc)
    return sh


def linha(slide, x1, y1, x2, y2, cor=BORDER, lw=1.0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                    Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = cor
    ln.line.width = Pt(lw)
    ln.shadow.inherit = False
    return ln


def icone(slide, nome, x, y, lado, cor="red"):
    return slide.shapes.add_picture(ic(nome, cor), Inches(x), Inches(y),
                                    Inches(lado), Inches(lado))


def cunha_direita(slide):
    """rtTriangle vermelho do canto direito, como nos layouts do CCU."""
    sh = forma(slide, MSO_SHAPE.RIGHT_TRIANGLE, 9.5634, 1.0748, 0.4343,
               4.5502, fill=RED)
    sh._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
    return sh


# --------------------------------------------------------------------------- #
# chrome padrao
# --------------------------------------------------------------------------- #
def rodape(slide, pagina=None, cor=INK):
    ins = (0.15, 0.15, 0.075, 0.075)
    texto(slide, 0.0, FOOT_Y, 5.3171, FOOT_H,
          [dict(runs=[(TURMA, {"bold": True}), (f"| {DECK}", {})],
                size=9, color=cor, font=F_FOOT)], insets=ins)
    texto(slide, 4.6827, FOOT_Y, 5.3171, FOOT_H,
          [dict(runs=COPY, size=9, color=cor, font=F_FOOT, align="r")],
          insets=ins)
    if pagina is not None:
        texto(slide, 4.6217, FOOT_Y, 0.2877, FOOT_H,
              [dict(runs=str(pagina), size=9, color=cor, font=F_FOOT,
                    align="c")])


def _layouts(prs):
    return {l.name: l for l in prs.slide_masters[0].slide_layouts}


def slide_branco(prs):
    return prs.slides.add_slide(_layouts(prs)["BLANK"])


def slide_conteudo(prs, titulo_secao, kicker, pagina):
    """Slide padrao: titulo de secao + kicker + cunha vermelha do layout."""
    slide = prs.slides.add_slide(_layouts(prs)["ONE_COLUMN_TEXT"])

    ph_titulo = None
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == 0:
            ph_titulo = ph
        else:
            ph._element.getparent().remove(ph._element)

    ph_titulo.left, ph_titulo.top = Inches(TITLE_BOX[0]), Inches(TITLE_BOX[1])
    ph_titulo.width = Inches(TITLE_BOX[2])
    ph_titulo.height = Inches(TITLE_BOX[3])
    tf = ph_titulo.text_frame
    tf.word_wrap = True
    set_insets(tf, 0.1, 0.1, 0.1, 0.1)
    limpar_tf(tf)
    par(tf, titulo_secao, size=24, color=RED, font=F_BODY)

    if kicker:
        texto(slide, *KICKER_BOX, [dict(runs=kicker, size=12, color=INK)])
    rodape(slide, pagina)
    return slide


def fonte_slide(slide, txt):
    texto(slide, BODY_L - 0.1, 4.93, BODY_W, 0.26,
          [dict(runs=f"Fonte: {txt}", size=8, color=GRAY, italic=True)])


def caixa_takeaway(slide, y, txt, h=0.47):
    """Takeaway box: filete vermelho + fundo cinza, no pe do exhibit."""
    rect(slide, BODY_L, y, BODY_W, h, fill=LIGHT)
    rect(slide, BODY_L, y, 0.055, h, fill=RED)
    texto(slide, BODY_L + 0.18, y, BODY_W - 0.32, h,
          [dict(runs=txt, size=11, color=INK, line=1.05)], anchor="m")


def chip_num(slide, x, y, lado, numero, size=14, dot_size=None, fill=RED,
             cor=LIGHT, font=F_NUM, align="c", largura=None):
    sh = forma(slide, MSO_SHAPE.RECTANGLE, x, y, largura or lado, lado,
               fill=fill)
    sh.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    runs = [(numero, {"size": size})]
    if dot_size:
        runs.append((".", {"size": dot_size}))
    par(sh.text_frame, runs, color=cor, font=font, align=align)
    return sh


# --------------------------------------------------------------------------- #
# slides especiais
# --------------------------------------------------------------------------- #
def moldura_escura(slide, assets, cunha=False):
    """Fundo #161616 full bleed + arte + logo, igual a capa oficial."""
    rect(slide, -0.0703, -0.0625, 10.1484, 5.7421, fill=INK, line=WHITE, lw=2.0)
    if cunha:
        cunha_direita(slide)
    pic = slide.shapes.add_picture(assets["arte-capa.png"], Inches(-0.0703),
                                   Inches(-0.0625), Inches(4.50),
                                   Inches(2.3403))
    pic.crop_top, pic.crop_bottom = 0.24866, 0.33529
    slide.shapes.add_picture(assets["logo-ccu.png"], Inches(0.4440),
                             Inches(4.9761), Inches(2.6107), Inches(0.3269))
    texto(slide, 4.6827, FOOT_Y, 5.3171, FOOT_H,
          [dict(runs=COPY, size=9, color=WHITE, font=F_FOOT, align="r")],
          insets=(0.15, 0.15, 0.075, 0.075))


def slide_capa(prs, assets):
    slide = slide_branco(prs)
    moldura_escura(slide, assets)
    texto(slide, 0.3359, 2.1200, 6.40, 1.8748,
          [dict(runs="APRESENTAÇÕES\nPARA CONSULTORIA", size=34, color=RED,
                font=F_DISP, line=0.94),
           dict(runs="Como estruturar, escrever e formatar slides",
                size=17, color=WHITE, font=F_DISP, before=8)], anchor="b")
    texto(slide, 0.3222, 4.0900, 5.0898, 0.5721,
          [dict(runs="Equipe de Conhecimento – CCU", size=13, color=WHITE),
           dict(runs="Prep4Consulting 2026.2 – Módulo complementar",
                size=13, color=WHITE)])
    return slide


def slide_agenda(prs, itens, pagina):
    slide = slide_branco(prs)
    texto(slide, 0.6675, 0.3894, 5.60, 0.6299,
          [dict(runs="VISÃO GERAL DO MÓDULO", size=28, color=RED,
                font=F_DISP, align="c")])
    y = 1.4224
    for num, label in itens:
        chip = chip_num(slide, 0.9843, y, 0.4553, num, size=14, dot_size=22,
                        largura=0.5273)
        set_insets(chip.text_frame, 0.05, 0.05, 0.02, 0.02)
        texto(slide, 1.6374, y, 6.60, 0.4553,
              [dict(runs=label, size=18, color=INK, font=F_DISP)], anchor="m")
        y += 0.7214
    rodape(slide, pagina)
    return slide


def slide_divisor(prs, num, linha1, linha2, pagina):
    slide = slide_branco(prs)
    chip = chip_num(slide, 7.6625, 1.7512, 1.3776, num, size=54, dot_size=54,
                    align="r")
    set_insets(chip.text_frame, 0.05, 0.05, 0.05, 0.05)
    texto(slide, 3.3333, 3.2353, 6.3225, 1.80,
          [dict(runs=[(linha1, {"color": RED}), ("\n" + linha2, {"color": INK})],
                size=40, font=F_DISP, align="r", line=0.98)])
    rodape(slide, pagina)
    return slide


def slide_encerramento(prs, assets):
    slide = slide_branco(prs)
    moldura_escura(slide, assets, cunha=True)
    texto(slide, 0.3359, 2.5000, 5.40, 1.7019,
          [dict(runs="AGRADECEMOS\nSUA PARTICIPAÇÃO!", size=30, color=RED,
                font=F_DISP, line=0.94),
           dict(runs="Agora é montar o seu deck.", size=17, color=WHITE,
                font=F_DISP, before=8)])
    texto(slide, 5.5700, 1.0600, 4.00, 0.40,
          [dict(runs="Ficou com alguma dúvida?", size=14, color=WHITE,
                bold=True)])
    contatos = [
        ("icone-email.png", "equipeccu@gmail.com", 1.7159),
        ("icone-instagram.png", "@ccuclubedeconsultoria", 2.5380),
        ("icone-site.png", "consultoriaunicamp.com", 3.3600),
        ("icone-linkedin.png", "Clube de Consultoria Universitário", 4.1821),
    ]
    for arquivo, label, y in contatos:
        rect(slide, 5.4584, y, 0.6730, 0.6730, fill=RED)
        slide.shapes.add_picture(assets[arquivo], Inches(5.5900),
                                 Inches(y + 0.131), Inches(0.41), Inches(0.41))
        texto(slide, 6.1300, y + 0.09, 3.38, 0.49,
              [dict(runs=label, size=13, color=WHITE, line=1.05)], anchor="m")
    return slide


# --------------------------------------------------------------------------- #
# blocos reutilizaveis
# --------------------------------------------------------------------------- #
def colunas_icone(slide, itens, y=1.72, h=2.40, lado=0.60, size_t=13,
                  size_d=10.5, gap=0.16):
    """Colunas iguais com icone, titulo e descricao centralizados."""
    n = len(itens)
    w = (BODY_W - gap * (n - 1)) / n
    for i, (nome, titulo, desc) in enumerate(itens):
        x = BODY_L + i * (w + gap)
        rect(slide, x, y, w, h, fill=WHITE, line=BORDER, lw=0.75)
        rect(slide, x, y, w, 0.055, fill=RED)
        icone(slide, nome, x + (w - lado) / 2, y + 0.26, lado)
        texto(slide, x + 0.06, y + 0.30 + lado + 0.06, w - 0.12, 0.50,
              [dict(runs=titulo, size=size_t, color=RED, bold=True,
                    align="c", line=1.0)])
        texto(slide, x + 0.07, y + 0.30 + lado + 0.56, w - 0.14,
              h - (0.30 + lado + 0.62),
              [dict(runs=desc, size=size_d, color=INK, align="c", line=1.14)])
    return w


def linhas_numeradas(slide, itens, x, y, w, dy=0.72, lado=0.30, size_t=12,
                     size_d=10):
    for i, (titulo, desc) in enumerate(itens, 1):
        yy = y + (i - 1) * dy
        chip = forma(slide, MSO_SHAPE.RECTANGLE, x, yy + 0.03, lado, lado,
                     fill=RED)
        chip.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        par(chip.text_frame, f"{i:02d}", size=10.5, color=WHITE, font=F_NUM,
            align="c")
        blocos = [dict(runs=titulo, size=size_t, color=INK, bold=True,
                       line=1.02)]
        if desc:
            blocos.append(dict(runs=desc, size=size_d, color=GRAY, line=1.1,
                               before=1.5))
        texto(slide, x + lado + 0.12, yy - 0.03, w - lado - 0.12, dy, blocos)


def harvey(slide, x, y, d, fracao):
    """Harvey ball: circulo com preenchimento proporcional."""
    forma(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=WHITE, line=RED, lw=1.25)
    if fracao >= 0.999:
        forma(slide, MSO_SHAPE.OVAL, x + 0.022, y + 0.022, d - 0.044,
              d - 0.044, fill=RED)
    elif fracao > 0.001:
        fim = int((16200000 + fracao * 21600000) % 21600000)
        forma(slide, MSO_SHAPE.PIE, x + 0.022, y + 0.022, d - 0.044, d - 0.044,
              fill=RED, adj=[("adj1", 16200000), ("adj2", fim)])


# --------------------------------------------------------------------------- #
# conteudo
# --------------------------------------------------------------------------- #
SEC = ["Estrutura antes do slide", "Anatomia do slide",
       "Dados e formatação", "Revisão e entrega"]


def pg_piramide(prs):
    s = slide_conteudo(prs, SEC[0],
                       "A resposta vem primeiro: a pirâmide organiza "
                       "recomendação, motivos e evidência", 3)
    cx = BODY_L + 2.65                      # centro do diagrama
    rect_txt(s, cx - 2.05, 1.70, 4.10, 0.58,
             "RECOMENDAÇÃO — a resposta em uma frase",
             fill=RED, size=11.5, bold=True)
    for i, rot in enumerate(["Motivo 1", "Motivo 2", "Motivo 3"]):
        bx = cx - 2.55 + i * 1.75
        linha(s, cx, 2.28, bx + 0.80, 2.60, cor=RED, lw=1.0)
        rect_txt(s, bx, 2.60, 1.60, 0.54, rot, fill=WHITE, line=RED, lw=1.25,
                 size=11.5, color=RED, bold=True)
        for j in range(3):
            ex = bx + 0.055 + j * 0.515
            linha(s, bx + 0.80, 3.14, ex + 0.22, 3.46, cor=BORDER, lw=0.75)
            rect(s, ex, 3.46, 0.44, 0.38, fill=LIGHT)
    texto(s, cx - 2.55, 3.92, 5.10, 0.28,
          [dict(runs="dados, análises e fatos que sustentam cada motivo",
                size=9.5, color=GRAY, align="c", italic=True)])
    linhas_numeradas(s, [
        ("Toda ideia resume o que está abaixo",
         "o nível de cima é síntese, não índice"),
        ("Ideias do mesmo nível são do mesmo grupo",
         "mesma natureza e nível de abstração"),
        ("Dentro do grupo existe uma ordem lógica",
         "por tempo, estrutura ou importância"),
    ], 6.06, 1.76, 3.42, dy=0.82, size_t=11.5, size_d=9.5)
    caixa_takeaway(s, 4.30,
                   "Se o cliente parar na primeira linha, ele já sabe o que "
                   "você recomenda — e por quais três motivos.")
    fonte_slide(s, "Barbara Minto, The Pyramid Principle; slideworks.io")


def pg_scqa(prs):
    s = slide_conteudo(prs, SEC[0],
                       "SCQA dá o contexto em quatro movimentos e entrega a "
                       "resposta antes da análise", 4)
    w = colunas_icone(s, [
        ("explore", "S · Situação",
         "O fato estável que o cliente aceita sem discussão"),
        ("warning", "C · Complicação",
         "O que mudou e criou tensão sobre essa situação"),
        ("help", "Q · Pergunta",
         "A dúvida que a complicação provoca no cliente"),
        ("task_alt", "A · Resposta",
         "Sua recomendação, dita antes de qualquer evidência"),
    ], y=1.72, h=2.38)
    for i in range(3):
        x = BODY_L + (i + 1) * (w + 0.16) - 0.135
        forma(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x, 2.32, 0.11, 0.17, fill=RED
              ).rotation = 90
    caixa_takeaway(s, 4.30,
                   "Regra prática: Situação e Complicação cabem em dois "
                   "parágrafos; a Resposta precisa caber em uma linha.")
    fonte_slide(s, "Minto, The Pyramid Principle; framework SCR/SCQA")


def pg_mece(prs):
    s = slide_conteudo(prs, SEC[0],
                       "Antes de abrir o PowerPoint: quebre o problema em MECE "
                       "e rascunhe o ghost deck", 5)
    passos = [
        ("account_tree", "1 · Quebrar o problema",
         "Árvore MECE: mutuamente exclusiva e coletivamente exaustiva — sem "
         "sobreposição entre ramos e sem lacuna no conjunto."),
        ("draw", "2 · Escrever o storyline",
         "Dot-dash: uma linha por slide, com o título-mensagem e os dois ou "
         "três pontos que vão sustentá-lo."),
        ("splitscreen", "3 · Montar o ghost deck",
         "Rascunho de cada slide com título, rótulo do dado e esboço do "
         "gráfico; a análise entra no lugar já reservado."),
    ]
    gw = (BODY_W - 0.26 * 2) / 3
    for i, (nome, titulo, desc) in enumerate(passos):
        x = BODY_L + i * (gw + 0.26)
        rect(s, x, 1.72, gw, 2.38, fill=WHITE, line=BORDER, lw=0.75)
        rect(s, x, 1.72, gw, 0.055, fill=RED)
        icone(s, nome, x + 0.22, 1.94, 0.54)
        texto(s, x + 0.20, 2.58, gw - 0.34, 0.32,
              [dict(runs=titulo, size=13, color=RED, bold=True)])
        texto(s, x + 0.20, 2.94, gw - 0.36, 1.05,
              [dict(runs=desc, size=10.5, color=INK, line=1.16)])
        if i < 2:
            forma(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x + gw + 0.075, 2.82, 0.11,
                  0.17, fill=RED).rotation = 90
    caixa_takeaway(s, 4.30,
                   "O erro mais caro de um deck não é o gráfico errado: é "
                   "abrir o PowerPoint antes de ter o storyline escrito.")
    fonte_slide(s, "strategyu.co (ghost deck); boardinfinity.com (storylining)")


def pg_anatomia(prs):
    s = slide_conteudo(prs, SEC[1],
                       "Todo slide de consultoria tem seis elementos fixos — e "
                       "nenhum deles é decoração", 7)
    mx, my, mw, mh = BODY_L, 1.68, 4.94, 3.05
    rect(s, mx, my, mw, mh, fill=WHITE, line=INK, lw=1.0)
    rect_txt(s, mx + 0.16, my + 0.18, mw - 0.90, 0.30,
             "título-mensagem em uma frase", fill=RED, size=9, bold=True,
             align="l")
    rect(s, mx + 0.16, my + 0.56, mw - 2.20, 0.15, fill=BORDER)
    rect(s, mx + 0.16, my + 0.84, mw - 2.26, 1.44, fill=LIGHT)
    for j, alt in enumerate([0.42, 0.68, 0.55, 0.92, 1.14]):
        rect(s, mx + 0.40 + j * 0.52, my + 2.20 - alt, 0.30, alt,
             fill=RED if j == 4 else GRAY)
    linha(s, mx + 2.44, my + 1.18, mx + 2.74, my + 1.18, cor=RED, lw=1.0)
    rect_txt(s, mx + 2.74, my + 0.97, 0.92, 0.42, "+14 p.p.", fill=WHITE,
             line=RED, lw=1.0, size=9, color=RED, bold=True)
    rect(s, mx + 0.16, my + 2.38, mw - 0.90, 0.26, fill=LIGHT)
    rect(s, mx + 0.16, my + 2.38, 0.04, 0.26, fill=RED)
    rect(s, mx + 0.16, my + 2.74, 1.90, 0.09, fill=BORDER)
    linha(s, mx, my + mh - 0.20, mx + mw, my + mh - 0.20, cor=BORDER, lw=0.75)
    for x, y, n in [(mx + mw - 0.58, my + 0.21, "1"),
                    (mx + mw - 0.58, my + 0.53, "2"),
                    (mx + mw - 0.58, my + 1.32, "3"),
                    (mx + 3.74, my + 1.07, "4"),
                    (mx + mw - 0.58, my + 2.37, "5"),
                    (mx + mw - 0.58, my + 2.70, "6")]:
        c = forma(s, MSO_SHAPE.OVAL, x, y, 0.23, 0.23, fill=INK)
        c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        par(c.text_frame, n, size=9, color=WHITE, font=F_NUM, align="c")
    linhas_numeradas(s, [
        ("Action title", "a mensagem do slide, em frase completa"),
        ("Kicker", "o recorte do dado: o que, onde e quando"),
        ("Exhibit", "um gráfico ou diagrama por slide, nunca dois"),
        ("Call-out", "a seta que aponta o número que importa"),
        ("Takeaway box", "o \u201ce daí?\u201d que fecha o raciocínio"),
        ("Fonte e rodapé", "origem do dado, numeração e copyright"),
    ], 5.62, 1.72, 3.86, dy=0.50, lado=0.25, size_t=11, size_d=9)
    fonte_slide(s, "theanalystacademy.com (takeaway box); slidescience.co")


def pg_action_title(prs):
    s = slide_conteudo(prs, SEC[1],
                       "Troque o título-rótulo pelo título-mensagem: ele "
                       "precisa responder \u201ce daí?\u201d sozinho", 8)
    cw = (BODY_W - 0.30) / 2
    casos = [
        (GRAY, "ANTES · título-rótulo", "cancel", "ink",
         "Análise de receita do 3º trimestre",
         "Diz do que o slide trata.",
         "O cliente precisa ler o slide inteiro para saber se a notícia é boa."),
        (RED, "DEPOIS · título-mensagem", "check_circle", "red",
         "A receita do 3º tri cresceu 14%, puxada por dois contratos "
         "corporativos",
         "Diz o que o slide significa.",
         "O cliente decide pelo título; o gráfico apenas comprova."),
    ]
    for i, (cor, rotulo, nome, cor_ic, titulo, legenda, nota) in enumerate(casos):
        x = BODY_L + i * (cw + 0.30)
        rect(s, x, 1.72, cw, 2.34, fill=WHITE, line=cor, lw=1.25)
        rect_txt(s, x, 1.72, cw, 0.30, rotulo, fill=cor, size=9.5, bold=True,
                 align="l")
        icone(s, nome, x + 0.16, 2.18, 0.30, cor_ic)
        texto(s, x + 0.54, 2.13, cw - 0.70, 0.76,
              [dict(runs=titulo, size=13.5, color=INK, bold=True, line=1.08)])
        texto(s, x + 0.16, 2.96, cw - 0.32, 0.26,
              [dict(runs=legenda, size=11, color=cor, bold=True)])
        texto(s, x + 0.16, 3.26, cw - 0.32, 0.70,
              [dict(runs=nota, size=10, color=GRAY, line=1.16)])
    regras = ["Frase completa, com verbo e número",
              "Uma única ideia por slide",
              "No máximo duas linhas de título",
              "Em sequência, os títulos contam a história"]
    rw = (BODY_W - 0.12 * 3) / 4
    for i, r in enumerate(regras):
        x = BODY_L + i * (rw + 0.12)
        rect(s, x, 4.20, rw, 0.54, fill=LIGHT)
        rect(s, x, 4.20, 0.045, 0.54, fill=RED)
        texto(s, x + 0.14, 4.20, rw - 0.24, 0.54,
              [dict(runs=r, size=9.5, color=INK, line=1.12)], anchor="m")
    fonte_slide(s, "slideworks.io; slidescience.co; highbridgeacademy.com")


def pg_graficos(prs):
    s = slide_conteudo(prs, SEC[2],
                       "Nomeie a comparação antes de escolher o gráfico — são "
                       "apenas cinco tipos", 10)
    colunas_icone(s, [
        ("donut_small", "Componente",
         "Parte sobre o total\n→ pizza ou barra 100%"),
        ("bar_chart", "Item",
         "Ranking entre itens\n→ barra horizontal ordenada"),
        ("trending_up", "Série temporal",
         "Evolução no tempo\n→ coluna ou linha"),
        ("stacked_bar_chart", "Distribuição",
         "Frequência por faixa\n→ histograma"),
        ("scatter_plot", "Correlação",
         "Duas variáveis juntas\n→ dispersão"),
    ], y=1.72, h=2.38, lado=0.54, size_t=12.5, size_d=10)
    caixa_takeaway(s, 4.30,
                   "A pergunta não é \u201cqual gráfico é mais bonito?\u201d, "
                   "e sim \u201cque comparação eu preciso provar?\u201d.")
    fonte_slide(s, "Gene Zelazny, Say It With Charts (as cinco comparações)")


def pg_elementos(prs):
    s = slide_conteudo(prs, SEC[2],
                       "Quatro recursos resolvem a maior parte dos exhibits de "
                       "um deck de consultoria", 11)
    qw, qh = (BODY_W - 0.24) / 2, 1.24
    quad = [(0, 0, "Harvey balls",
             "Comparar opções em critérios qualitativos, na mesma escala."),
            (1, 0, "Waterfall",
             "Explicar de onde veio a variação entre dois pontos."),
            (0, 1, "Matriz 2×2",
             "Posicionar alternativas e priorizar o que entra no plano."),
            (1, 1, "Status por frente",
             "Mostrar andamento sem precisar de texto corrido.")]
    for col, row, titulo, desc in quad:
        x = BODY_L + col * (qw + 0.24)
        y = 1.68 + row * (qh + 0.20)
        rect(s, x, y, qw, qh, fill=WHITE, line=BORDER, lw=0.75)
        rect(s, x, y, 0.045, qh, fill=RED)
        texto(s, x + 0.16, y + 0.10, qw * 0.50, 0.26,
              [dict(runs=titulo, size=12.5, color=RED, bold=True)])
        texto(s, x + 0.17, y + 0.39, qw * 0.50, 0.76,
              [dict(runs=desc, size=10, color=INK, line=1.16)])
        gx, gy = x + qw * 0.57, y + 0.14
        if titulo == "Harvey balls":
            for j, f in enumerate([0.0, 0.25, 0.5, 0.75, 1.0]):
                harvey(s, gx + j * 0.32, gy + 0.22, 0.25, f)
            texto(s, gx, gy + 0.56, 1.62, 0.22,
                  [dict(runs="pior   →   melhor", size=8, color=GRAY,
                        align="c")])
        elif titulo == "Waterfall":
            base = gy + 0.86
            # (deslocamento x, altura, base acumulada, cor)
            for bx, bh, off, cor in [(0.00, 0.44, 0.00, GRAY),
                                     (0.34, 0.14, 0.44, RED),
                                     (0.68, 0.10, 0.58, RED),
                                     (1.02, 0.18, 0.50, RED_DARK),
                                     (1.36, 0.50, 0.00, GRAY)]:
                rect(s, gx + bx, base - off - bh, 0.24, bh, fill=cor)
            linha(s, gx - 0.04, base, gx + 1.66, base, cor=GRAY, lw=0.75)
        elif titulo == "Matriz 2×2":
            for a in range(2):
                for b in range(2):
                    rect(s, gx + 0.34 + b * 0.54, gy + 0.10 + a * 0.33,
                         0.54, 0.33,
                         fill=RED if (a, b) == (0, 1) else LIGHT,
                         line=BORDER, lw=0.6)
            linha(s, gx + 0.32, gy + 0.08, gx + 0.32, gy + 0.78, cor=INK)
            linha(s, gx + 0.30, gy + 0.78, gx + 1.46, gy + 0.78, cor=INK)
        else:
            for j, cor in enumerate([RED, GRAY, BORDER]):
                forma(s, MSO_SHAPE.OVAL, gx + 0.34, gy + 0.12 + j * 0.24,
                      0.17, 0.17, fill=cor)
                rect(s, gx + 0.58, gy + 0.165 + j * 0.24, 0.82 - j * 0.20,
                     0.07, fill=BORDER)
    caixa_takeaway(s, 4.38,
                   "Escolha o recurso pela pergunta que ele responde — não "
                   "pelo efeito visual que produz.", h=0.42)
    fonte_slide(s, "umbrex.com; managementconsulted.com")


def pg_formatacao(prs):
    s = slide_conteudo(prs, SEC[2],
                       "Formatação não é estética: ela reduz o esforço de "
                       "leitura do cliente", 12)
    colunas_icone(s, [
        ("grid_on", "Grade e alinhamento",
         "Margens fixas; nada \u201cquase\u201d alinhado."),
        ("palette", "Cor com função",
         "Uma cor de destaque para o número que importa."),
        ("format_size", "Tipografia contida",
         "Uma família sem serifa e mínimo de 10 pt."),
        ("crop_free", "Espaço em branco",
         "O vazio é o que faz o dado aparecer."),
    ], y=1.70, h=2.10, lado=0.52, size_t=12.5, size_d=10)
    texto(s, BODY_L, 3.98, 4.00, 0.26,
          [dict(runs="Paleta deste template", size=10.5, color=INK,
                bold=True)])
    paleta = [("#C1272D", RED, "destaque"), ("#161616", INK, "texto"),
              ("#F3F3F3", LIGHT, "fundo de apoio"),
              ("#888888", GRAY, "texto secundário")]
    for i, (hexa, cor, uso) in enumerate(paleta):
        x = BODY_L + i * 2.26
        rect(s, x, 4.28, 0.40, 0.40, fill=cor,
             line=BORDER if cor is LIGHT else None)
        texto(s, x + 0.48, 4.26, 1.74, 0.44,
              [dict(runs=hexa, size=10, color=INK, bold=True, font=F_NUM),
               dict(runs=uso, size=9, color=GRAY)])
    fonte_slide(s, "umbrex.com (design de charts); blog.drlambda.ai")


def pg_checklist(prs):
    s = slide_conteudo(prs, SEC[3],
                       "Dez verificações antes de enviar: se alguma falhar, o "
                       "slide volta para a mesa", 14)
    itens = [
        "Lendo só os títulos, a história se sustenta",
        "Cada slide carrega uma única mensagem",
        "Todo exhibit tem um \u201ce daí?\u201d explícito",
        "Todo número tem fonte e data",
        "Os grupos de argumentos são MECE",
        "Nenhum gráfico depende de legenda para ser lido",
        "Margens e alinhamentos iguais em todos os slides",
        "Nenhuma sigla sem definição na primeira aparição",
        "A recomendação aparece antes da terceira página",
        "O deck se explica sozinho, sem o apresentador",
    ]
    colw = (BODY_W - 0.32) / 2
    for i, it in enumerate(itens):
        col, row = i // 5, i % 5
        x = BODY_L + col * (colw + 0.32)
        y = 1.72 + row * 0.51
        rect(s, x, y, colw, 0.43, fill=LIGHT if row % 2 == 0 else WHITE,
             line=None if row % 2 == 0 else BORDER, lw=0.6)
        icone(s, "check_circle", x + 0.12, y + 0.10, 0.23)
        texto(s, x + 0.42, y, colw - 0.80, 0.43,
              [dict(runs=it, size=10.5, color=INK, line=1.1)], anchor="m")
        texto(s, x + colw - 0.44, y, 0.36, 0.43,
              [dict(runs=f"{i + 1:02d}", size=10, color=GRAY, font=F_NUM,
                    align="r")], anchor="m")
    caixa_takeaway(s, 4.36,
                   "Rode a lista com alguém que não participou da análise: o "
                   "que precisar de explicação, reescreva.", h=0.44)


def pg_erros(prs):
    s = slide_conteudo(prs, SEC[3],
                       "Seis erros aparecem em quase toda primeira versão de "
                       "um deck", 15)
    erros = [
        ("Título-rótulo", "reescreva como frase com verbo e número."),
        ("Slide-dump", "uma mensagem por slide; o excedente vai para o anexo."),
        ("Gráfico sem conclusão", "acrescente call-out e takeaway box."),
        ("Dado sem fonte", "fonte e data no pé do exhibit, sempre."),
        ("Grupos sobrepostos", "teste MECE: sem sobreposição e sem lacuna."),
        ("Formatação inconsistente", "use o template; não mova caixas no olho."),
    ]
    ew = (BODY_W - 0.22 * 2) / 3
    for i, (erro, correcao) in enumerate(erros):
        col, row = i % 3, i // 3
        x = BODY_L + col * (ew + 0.22)
        y = 1.72 + row * 1.30
        rect(s, x, y, ew, 1.14, fill=WHITE, line=BORDER, lw=0.75)
        icone(s, "cancel", x + 0.14, y + 0.15, 0.25)
        texto(s, x + 0.45, y + 0.12, ew - 0.58, 0.32,
              [dict(runs=erro, size=12, color=RED, bold=True, line=1.0)])
        texto(s, x + 0.15, y + 0.52, ew - 0.30, 0.56,
              [dict(runs=[("Corrija: ", {"bold": True}), (correcao, {})],
                    size=10, color=INK, line=1.16)])
    caixa_takeaway(s, 4.40,
                   "Os seis são erros de disciplina, não de talento: revisão "
                   "com checklist resolve.", h=0.44)


def pg_takeaways(prs):
    s = slide_conteudo(prs, "Key takeaways",
                       "O que levar deste módulo para o próximo deck que você "
                       "montar", 16)
    tk = [
        ("A história vem antes do slide",
         "Storyline em dez linhas e ghost deck no papel; o PowerPoint é a "
         "última etapa, não a primeira."),
        ("Resposta primeiro",
         "Pirâmide e SCQA entregam a recomendação na abertura e deixam a "
         "evidência para depois."),
        ("Um slide, uma mensagem",
         "Um exhibit, um call-out e um \u201ce daí?\u201d por slide — o resto "
         "vai para o anexo."),
        ("Formatação é disciplina",
         "Grade, uma cor de destaque, fonte em todo número e checklist antes "
         "de enviar."),
    ]
    for i, (titulo, desc) in enumerate(tk):
        y = 1.70 + i * 0.76
        chip_num(s, BODY_L, y, 0.56, f"{i + 1:02d}", size=15)
        texto(s, BODY_L + 0.72, y - 0.02, BODY_W - 0.80, 0.62,
              [dict(runs=titulo, size=13.5, color=RED, bold=True, line=1.0),
               dict(runs=desc, size=10.5, color=INK, line=1.16, before=2)])
    caixa_takeaway(s, 4.78,
                   "Um bom deck de consultoria é um documento que decide, "
                   "não um resumo do que foi analisado.", h=0.40)


def pg_fontes(prs):
    s = slide_conteudo(prs, "Leitura recomendada",
                       "Material de referência em inglês, sintetizado e "
                       "reescrito para este módulo", 17)
    refs = [
        ("Livro", "Barbara Minto — The Pyramid Principle",
         "estrutura de argumento, SCQA e agrupamento MECE", None),
        ("Livro", "Gene Zelazny — Say It With Charts",
         "as cinco comparações e a escolha do gráfico", None),
        ("Artigo", "slideworks.io — How to write action titles like McKinsey",
         "títulos-mensagem com exemplos de antes e depois",
         "https://slideworks.io/resources/how-to-write-action-titles-like-mckinsey"),
        ("Artigo", "slidescience.co — What is the Pyramid Principle",
         "aplicação da pirâmide no nível do slide",
         "https://slidescience.co/pyramid-principle/"),
        ("Artigo", "strategyu.co — The ghost deck",
         "como rascunhar o deck antes de ter a análise",
         "https://strategyu.co/ghost-deck/"),
        ("Artigo", "theanalystacademy.com — Takeaway boxes",
         "quando usar e quando evitar a caixa de conclusão",
         "https://www.theanalystacademy.com/takeaway-boxes-when-to-use/"),
        ("Artigo", "umbrex.com — Design principles for quantitative charts",
         "padrões de formatação de gráficos em consultoria",
         "https://umbrex.com/resources/the-busy-consultants-guide-to-quantitative-charts/design-principles-for-mckinsey-quantitative-charts/"),
        ("Artigo", "managementconsulted.com — Consulting slide deck",
         "convenções de estrutura e erros recorrentes",
         "https://managementconsulted.com/consulting-slide-deck/"),
    ]
    for i, (tipo, titulo, nota, link) in enumerate(refs):
        y = 1.64 + i * 0.395
        if i % 2 == 0:
            rect(s, BODY_L, y, BODY_W, 0.37, fill=LIGHT)
        rect_txt(s, BODY_L + 0.08, y + 0.075, 0.54, 0.22, tipo.upper(),
                 fill=RED if tipo == "Livro" else INK, size=7)
        ov = {"bold": True, "color": INK}
        if link:
            ov["link"] = link
        texto(s, BODY_L + 0.70, y, BODY_W - 0.80, 0.37,
              [dict(runs=[(titulo, ov),
                          (f"  ·  {nota}", {"color": GRAY, "size": 9.5})],
                    size=10.5)], anchor="m")
    texto(s, BODY_L, 4.84, BODY_W, 0.30,
          [dict(runs="Conteúdo de terceiros citado para estudo; os textos "
                     "foram parafraseados e resumidos pela equipe do CCU.",
                size=8, color=GRAY, italic=True)])


# --------------------------------------------------------------------------- #
# montagem
# --------------------------------------------------------------------------- #
def construir():
    assets = extrair_assets()
    shutil.copyfile(TEMPLATE, OUT)
    prs = Presentation(OUT)
    limpar_slides(prs)

    slide_capa(prs, assets)
    slide_agenda(prs, [("01", SEC[0]), ("02", SEC[1]),
                       ("03", SEC[2]), ("04", SEC[3])], 1)

    slide_divisor(prs, "01", "Estrutura antes ", "do slide", 2)
    pg_piramide(prs)
    pg_scqa(prs)
    pg_mece(prs)

    slide_divisor(prs, "02", "Anatomia ", "do slide", 6)
    pg_anatomia(prs)
    pg_action_title(prs)

    slide_divisor(prs, "03", "Dados e ", "formatação", 9)
    pg_graficos(prs)
    pg_elementos(prs)
    pg_formatacao(prs)

    slide_divisor(prs, "04", "Revisão ", "e entrega", 13)
    pg_checklist(prs)
    pg_erros(prs)
    pg_takeaways(prs)
    pg_fontes(prs)

    slide_encerramento(prs, assets)

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print("gerado:", construir())
